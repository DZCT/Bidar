"""
Tests for the Universal Music Downloader + Allow whitelist + auto-detect.
Runs without a live Telegram session — uses mocks only.

Run:
    cd /app/bidar && python tests/test_music_universal.py
"""
from __future__ import annotations

import os
import sys
import copy
import io
import base64
import asyncio
import tempfile
import unittest
from pathlib import Path
from unittest.mock import MagicMock, AsyncMock, patch

# Stub env vars BEFORE importing bidar (which reads them at import time)
os.environ.setdefault("API_ID", "12345")
os.environ.setdefault("API_HASH", "x")
os.environ.setdefault("PHONE", "x")

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
import bidar  # noqa: E402


# ────────────────────────────────────────────────────────────────────
# 1. URL Detection
# ────────────────────────────────────────────────────────────────────
class TestUrlDetection(unittest.TestCase):
    def test_soundcloud_url(self):
        got = bidar._detect_music_url("Check https://soundcloud.com/artist/song-title cool")
        self.assertIsNotNone(got)
        self.assertEqual(got[0], "soundcloud")
        self.assertFalse(got[2])

    def test_youtube_music_only(self):
        """Only music.youtube.com is treated as music — regular video links are ignored."""
        # Music: detected as YouTube Music
        got = bidar._detect_music_url("https://music.youtube.com/watch?v=dQw4w9WgXcQ")
        self.assertIsNotNone(got)
        self.assertEqual(got[0], "youtube")
        got = bidar._detect_music_url("https://music.youtube.com/playlist?list=PLabc123")
        self.assertIsNotNone(got)
        self.assertEqual(got[0], "youtube")
        # Regular YouTube video / short links must NOT be auto-handled
        for url in [
            "https://www.youtube.com/watch?v=dQw4w9WgXcQ",
            "https://youtu.be/dQw4w9WgXcQ",
            "https://m.youtube.com/watch?v=dQw4w9WgXcQ",
            "https://youtube.com/shorts/abc",
        ]:
            self.assertIsNone(bidar._detect_music_url(url), msg=f"should be ignored: {url}")

    def test_spotify_url(self):
        got = bidar._detect_music_url(
            "track https://open.spotify.com/track/4uLU6hMCjMI75M1A2tKUQC enjoy"
        )
        self.assertIsNotNone(got)
        self.assertEqual(got[0], "spotify")
        self.assertTrue(got[2])

    def test_spotify_intl_url(self):
        got = bidar._detect_music_url(
            "https://open.spotify.com/intl-fa/track/4uLU6hMCjMI75M1A2tKUQC"
        )
        self.assertEqual(got[0], "spotify")

    def test_deezer_url(self):
        got = bidar._detect_music_url("https://www.deezer.com/track/123456789")
        self.assertEqual(got[0], "deezer")
        self.assertTrue(got[2])

    def test_apple_music_url(self):
        got = bidar._detect_music_url(
            "https://music.apple.com/us/song/blinding-lights/1499378108"
        )
        self.assertEqual(got[0], "apple")
        self.assertTrue(got[2])

    def test_tidal_url(self):
        got = bidar._detect_music_url("https://tidal.com/browse/track/12345")
        self.assertEqual(got[0], "tidal")
        self.assertTrue(got[2])

    def test_bandcamp_url(self):
        got = bidar._detect_music_url("https://artist.bandcamp.com/track/song-name")
        self.assertEqual(got[0], "bandcamp")
        self.assertFalse(got[2])

    def test_mixcloud_url(self):
        got = bidar._detect_music_url("https://www.mixcloud.com/user/track-name/")
        self.assertEqual(got[0], "mixcloud")

    def test_no_music_url(self):
        for noise in [
            "",
            "just a normal message",
            "https://google.com",
            "https://twitter.com/foo",
            "https://example.com/song",
            "soundcloud is great",  # text mention only, no URL
        ]:
            self.assertIsNone(
                bidar._detect_music_url(noise),
                msg=f"False positive for: {noise!r}",
            )


# ────────────────────────────────────────────────────────────────────
# 2. DRM Metadata extraction (mock urllib.request)
# ────────────────────────────────────────────────────────────────────
class TestDrmMetadata(unittest.TestCase):
    def _mock_resp(self, html: str):
        m = MagicMock()
        m.__enter__ = lambda s: s
        m.__exit__ = lambda *a: None
        m.read.return_value = html.encode("utf-8")
        return m

    def test_spotify(self):
        html = (
            '<meta property="og:title" content="Blinding Lights">'
            '<meta property="og:description" '
            'content="Listen to Blinding Lights on Spotify. The Weeknd · Song · 2019">'
        )
        with patch("urllib.request.urlopen", return_value=self._mock_resp(html)):
            q = bidar._fetch_drm_metadata("https://open.spotify.com/track/abc", "spotify")
        self.assertIsNotNone(q)
        self.assertIn("Blinding Lights", q)
        self.assertIn("The Weeknd", q)

    def test_deezer(self):
        html = (
            '<meta property="og:title" content="Shape of You">'
            '<meta property="og:description" content="Listen to Shape of You by Ed Sheeran on Deezer">'
        )
        with patch("urllib.request.urlopen", return_value=self._mock_resp(html)):
            q = bidar._fetch_drm_metadata("https://deezer.com/track/1", "deezer")
        self.assertIn("Shape of You", q)
        self.assertIn("Ed Sheeran", q)

    def test_apple(self):
        html = (
            '<meta property="og:title" content="Bohemian Rhapsody">'
            '<meta property="og:description" content="Song · Queen · 1975">'
        )
        with patch("urllib.request.urlopen", return_value=self._mock_resp(html)):
            q = bidar._fetch_drm_metadata(
                "https://music.apple.com/us/song/bohemian-rhapsody/1", "apple"
            )
        self.assertIn("Bohemian Rhapsody", q)
        self.assertIn("Queen", q)

    def test_no_metadata(self):
        with patch("urllib.request.urlopen", return_value=self._mock_resp("<html></html>")):
            q = bidar._fetch_drm_metadata("https://open.spotify.com/track/x", "spotify")
        self.assertIsNone(q)

    def test_network_error(self):
        with patch("urllib.request.urlopen", side_effect=Exception("network")):
            q = bidar._fetch_drm_metadata("https://open.spotify.com/track/x", "spotify")
        self.assertIsNone(q)


# ────────────────────────────────────────────────────────────────────
# 3. Allow whitelist
# ────────────────────────────────────────────────────────────────────
class TestAllowList(unittest.TestCase):
    def setUp(self):
        self.tmpdir = tempfile.mkdtemp(prefix="bidar_test_")
        bidar.CONFIG_FILE = Path(self.tmpdir) / "bidar_config.json"
        bidar.config = copy.deepcopy(bidar._DEFAULT_CONFIG)
        bidar.save_config()

    def test_default_is_empty(self):
        self.assertEqual(bidar.config["allowed_groups"], [])

    def test_add_and_persist(self):
        bidar.config["allowed_groups"].append(-100123)
        bidar.save_config()
        bidar.load_config()
        self.assertIn(-100123, bidar.config["allowed_groups"])

    def test_music_enabled_default(self):
        self.assertTrue(bidar.config["music_enabled"])


# ────────────────────────────────────────────────────────────────────
# 4. End-to-end auto-detect flow
# ────────────────────────────────────────────────────────────────────
class TestAutoDetectFlow(unittest.IsolatedAsyncioTestCase):
    def _mk_event(self, *, text, chat_id, is_private):
        ev = MagicMock()
        ev.raw_text = text
        ev.chat_id = chat_id
        ev.is_private = is_private
        ev.message = MagicMock()
        ev.message.id = 42
        return ev

    async def test_pv_always_allowed(self):
        bidar.config = copy.deepcopy(bidar._DEFAULT_CONFIG)
        bidar.config["music_enabled"] = True
        bidar.config["allowed_groups"] = []
        bidar.YTDLP_OK = True
        bidar._music_recent.clear()

        called = {}

        async def fake_dl(event, url, platform, is_drm, **kw):
            called["platform"] = platform
            called["is_drm"] = is_drm
            called["url"] = url

        ev = self._mk_event(
            text="Check https://open.spotify.com/track/abc",
            chat_id=999, is_private=True,
        )
        sender = MagicMock(id=111, bot=False)
        with patch.object(bidar, "_music_download_and_send", side_effect=fake_dl):
            await bidar._maybe_handle_music_link(ev, sender)
        self.assertEqual(called.get("platform"), "spotify")
        self.assertTrue(called.get("is_drm"))

    async def test_group_not_whitelisted_skipped(self):
        bidar.config = copy.deepcopy(bidar._DEFAULT_CONFIG)
        bidar.config["music_enabled"] = True
        bidar.config["allowed_groups"] = [12345]
        bidar.YTDLP_OK = True
        bidar._music_recent.clear()
        called = {"n": 0}

        async def fake_dl(*a, **kw):
            called["n"] += 1

        ev = self._mk_event(
            text="https://soundcloud.com/foo/bar",
            chat_id=-100999, is_private=False,
        )
        sender = MagicMock(id=111, bot=False)
        with patch.object(bidar, "_music_download_and_send", side_effect=fake_dl):
            await bidar._maybe_handle_music_link(ev, sender)
        self.assertEqual(called["n"], 0)

    async def test_group_whitelisted_allowed(self):
        bidar.config = copy.deepcopy(bidar._DEFAULT_CONFIG)
        bidar.config["music_enabled"] = True
        bidar.config["allowed_groups"] = [-100999]
        bidar.YTDLP_OK = True
        bidar._music_recent.clear()
        called = {"url": None}

        async def fake_dl(event, url, platform, is_drm, **kw):
            called["url"] = url

        ev = self._mk_event(
            text="https://soundcloud.com/foo/bar",
            chat_id=-100999, is_private=False,
        )
        sender = MagicMock(id=111, bot=False)
        with patch.object(bidar, "_music_download_and_send", side_effect=fake_dl):
            await bidar._maybe_handle_music_link(ev, sender)
        self.assertIn("soundcloud.com/foo/bar", called["url"] or "")

    async def test_music_disabled_skips(self):
        bidar.config = copy.deepcopy(bidar._DEFAULT_CONFIG)
        bidar.config["music_enabled"] = False
        bidar.YTDLP_OK = True
        bidar._music_recent.clear()
        called = {"n": 0}

        async def fake_dl(*a, **kw):
            called["n"] += 1

        ev = self._mk_event(
            text="https://open.spotify.com/track/abc",
            chat_id=999, is_private=True,
        )
        sender = MagicMock(id=111, bot=False)
        with patch.object(bidar, "_music_download_and_send", side_effect=fake_dl):
            await bidar._maybe_handle_music_link(ev, sender)
        self.assertEqual(called["n"], 0)

    async def test_dedup_within_5min(self):
        bidar.config = copy.deepcopy(bidar._DEFAULT_CONFIG)
        bidar.config["music_enabled"] = True
        bidar.YTDLP_OK = True
        bidar._music_recent.clear()
        count = {"n": 0}

        async def fake_dl(*a, **kw):
            count["n"] += 1
            return True  # success keeps the dedup entry

        ev = self._mk_event(
            text="https://open.spotify.com/track/dedupe",
            chat_id=9876, is_private=True,
        )
        sender = MagicMock(id=111, bot=False)
        with patch.object(bidar, "_music_download_and_send", side_effect=fake_dl):
            await bidar._maybe_handle_music_link(ev, sender)
            await bidar._maybe_handle_music_link(ev, sender)  # immediate retry
        self.assertEqual(count["n"], 1)


# ────────────────────────────────────────────────────────────────────
# 5. i18n: ensure all new keys are present for both languages
# ────────────────────────────────────────────────────────────────────
class TestI18N(unittest.TestCase):
    def test_new_keys_present(self):
        keys = [
            "music_downloading", "music_resolving", "music_uploading",
            "music_failed", "music_meta_failed", "music_caption", "music_set",
            "allow_help", "allow_added", "allow_exists", "allow_removed",
            "allow_notfound", "allow_invalid", "allow_list_empty",
            "allow_list_title", "allow_cleared", "allow_here_pv",
        ]
        for k in keys:
            entry = bidar.I18N.get(k, {})
            self.assertTrue(entry.get("en"), msg=f"missing en for {k}")
            self.assertTrue(entry.get("fa"), msg=f"missing fa for {k}")

    def test_help_mentions_new_commands(self):
        en = bidar.I18N["help_full"]["en"]
        fa = bidar.I18N["help_full"]["fa"]
        self.assertIn("music on|off", en)
        self.assertIn("allow help", en)
        self.assertIn("music on|off", fa)
        self.assertIn("allow help", fa)


# ────────────────────────────────────────────────────────────────────
# 6. v1.9.1 fixes: shortlinks, chat-ID variants, outgoing links, dedup retry
# ────────────────────────────────────────────────────────────────────
class TestShortlinkDetection(unittest.TestCase):
    def test_on_soundcloud_shortlink_detected(self):
        got = bidar._detect_music_url("https://on.soundcloud.com/aquUJPmourEr5c6Lul")
        self.assertIsNotNone(got)
        self.assertEqual(got[0], "soundcloud")
        self.assertFalse(got[2])

    def test_spotify_link_short_detected_as_drm(self):
        got = bidar._detect_music_url("listen https://spotify.link/AbC123xyz now")
        self.assertIsNotNone(got)
        self.assertEqual(got[0], "spotify")
        self.assertTrue(got[2])

    def test_deezer_page_link_detected_as_drm(self):
        got = bidar._detect_music_url("https://deezer.page.link/XyZ987")
        self.assertIsNotNone(got)
        self.assertEqual(got[0], "deezer")
        self.assertTrue(got[2])

    def test_shortlink_re_matches_only_drm_shorteners(self):
        self.assertTrue(bidar._SHORTLINK_RE.match("https://spotify.link/abc"))
        self.assertTrue(bidar._SHORTLINK_RE.match("https://deezer.page.link/x"))
        self.assertIsNone(bidar._SHORTLINK_RE.match("https://on.soundcloud.com/x"))
        self.assertIsNone(bidar._SHORTLINK_RE.match("https://open.spotify.com/track/x"))


class TestChatIdVariants(unittest.TestCase):
    def test_marked_supergroup_includes_raw(self):
        v = bidar._chat_id_variants(-1002453861964)
        self.assertIn(2453861964, v)
        self.assertIn(-1002453861964, v)

    def test_raw_positive_includes_marked(self):
        v = bidar._chat_id_variants(2453861964)
        self.assertIn(-1002453861964, v)

    def test_legacy_group_includes_positive(self):
        v = bidar._chat_id_variants(-456789)
        self.assertIn(456789, v)
        self.assertIn(-100456789, v)

    def test_invalid_returns_empty(self):
        self.assertEqual(bidar._chat_id_variants("abc"), set())

    def test_is_chat_allowed_raw_id_matches_marked_chat(self):
        bidar.config = copy.deepcopy(bidar._DEFAULT_CONFIG)
        bidar.config["allowed_groups"] = [2453861964]  # raw `.id`-style entry
        self.assertTrue(bidar._is_chat_allowed(-1002453861964))

    def test_is_chat_allowed_marked_entry_matches(self):
        bidar.config = copy.deepcopy(bidar._DEFAULT_CONFIG)
        bidar.config["allowed_groups"] = [-1002453861964]
        self.assertTrue(bidar._is_chat_allowed(-1002453861964))

    def test_is_chat_allowed_rejects_other_chat(self):
        bidar.config = copy.deepcopy(bidar._DEFAULT_CONFIG)
        bidar.config["allowed_groups"] = [111222333]
        self.assertFalse(bidar._is_chat_allowed(-1009998887776))

    def test_whitelist_contains_and_without(self):
        lst = [2453861964, -100555]
        self.assertTrue(bidar._whitelist_contains(-1002453861964, lst))
        remaining = bidar._whitelist_without(-1002453861964, lst)
        self.assertEqual(remaining, [-100555])


class TestAutoDetectV191(unittest.IsolatedAsyncioTestCase):
    def _mk_event(self, *, text, chat_id, is_private):
        ev = MagicMock()
        ev.raw_text = text
        ev.chat_id = chat_id
        ev.is_private = is_private
        ev.message = MagicMock()
        ev.message.id = 42
        return ev

    async def test_group_raw_id_whitelist_triggers(self):
        """Group stored with raw positive `.id` output must still trigger."""
        bidar.config = copy.deepcopy(bidar._DEFAULT_CONFIG)
        bidar.config["music_enabled"] = True
        bidar.config["allowed_groups"] = [2453861964]
        bidar.YTDLP_OK = True
        bidar._music_recent.clear()
        called = {"n": 0}

        async def fake_dl(*a, **kw):
            called["n"] += 1
            return True

        ev = self._mk_event(text="https://soundcloud.com/foo/bar",
                            chat_id=-1002453861964, is_private=False)
        with patch.object(bidar, "_music_download_and_send", side_effect=fake_dl):
            await bidar._maybe_handle_music_link(ev, MagicMock(id=1, bot=False))
        self.assertEqual(called["n"], 1)

    async def test_outgoing_owner_message_sender_none(self):
        """Owner's own (outgoing) link → sender=None must still work in PV."""
        bidar.config = copy.deepcopy(bidar._DEFAULT_CONFIG)
        bidar.config["music_enabled"] = True
        bidar.YTDLP_OK = True
        bidar._music_recent.clear()
        called = {"kw": None}

        async def fake_dl(event, url, platform, is_drm, **kw):
            called["kw"] = kw
            return True

        ev = self._mk_event(text="https://on.soundcloud.com/abc123",
                            chat_id=555, is_private=True)
        with patch.object(bidar, "_music_download_and_send", side_effect=fake_dl):
            await bidar._maybe_handle_music_link(ev, None)
        self.assertIsNotNone(called["kw"])
        self.assertTrue(called["kw"].get("force_reply"))

    async def test_failed_download_clears_dedup_for_retry(self):
        """A failed download must NOT block an immediate retry of the same link."""
        bidar.config = copy.deepcopy(bidar._DEFAULT_CONFIG)
        bidar.config["music_enabled"] = True
        bidar.YTDLP_OK = True
        bidar._music_recent.clear()
        count = {"n": 0}

        async def fake_dl(*a, **kw):
            count["n"] += 1
            return False  # simulate failure

        ev = self._mk_event(text="https://soundcloud.com/retry/me",
                            chat_id=777, is_private=True)
        sender = MagicMock(id=111, bot=False)
        with patch.object(bidar, "_music_download_and_send", side_effect=fake_dl):
            await bidar._maybe_handle_music_link(ev, sender)
            await bidar._maybe_handle_music_link(ev, sender)  # retry after failure
        self.assertEqual(count["n"], 2)


# ────────────────────────────────────────────────────────────────────
# 7. v1.9.2: metadata APIs (iTunes/Deezer), query cleaner, YT oEmbed fallback
# ────────────────────────────────────────────────────────────────────
class TestCleanQuery(unittest.TestCase):
    def test_unescape_and_nbsp(self):
        self.assertEqual(bidar._clean_query("Let&#39;s Go on Apple\xa0Music"),
                         "Let's Go on Apple Music")

    def test_strips_official_video(self):
        self.assertEqual(bidar._clean_query("Artist - Song (Official Video)"), "Artist - Song")
        self.assertEqual(bidar._clean_query("Artist - Song [Official Lyric Video]"), "Artist - Song")

    def test_strips_single_suffix(self):
        self.assertEqual(bidar._clean_query("The Weeknd - Blinding Lights - Single"),
                         "The Weeknd - Blinding Lights")


class TestPlatformLookups(unittest.TestCase):
    def test_apple_lookup_uses_i_param(self):
        calls = []

        def fake_json(url, timeout=15):
            calls.append(url)
            return {"results": [{"artistName": "Joost", "trackName": "Europapa"}]}

        with patch.object(bidar, "_http_json", side_effect=fake_json):
            q = bidar._apple_lookup("https://music.apple.com/us/album/europapa/1732041797?i=1732041802")
        self.assertEqual(q, "Joost - Europapa")
        self.assertIn("id=1732041802", calls[0])
        self.assertIn("country=us", calls[0])

    def test_apple_lookup_song_path_id(self):
        def fake_json(url, timeout=15):
            return {"results": [{"artistName": "Joost", "trackName": "Europapa"}]}

        with patch.object(bidar, "_http_json", side_effect=fake_json):
            q = bidar._apple_lookup("https://music.apple.com/us/song/europapa/1732041802")
        self.assertEqual(q, "Joost - Europapa")

    def test_apple_lookup_dead_link_returns_none(self):
        with patch.object(bidar, "_http_json", return_value={"resultCount": 0, "results": []}):
            q = bidar._apple_lookup("https://music.apple.com/us/album/x/1499385311?i=1499385316")
        self.assertIsNone(q)

    def test_deezer_lookup(self):
        def fake_json(url, timeout=15):
            assert "api.deezer.com/track/3135556" in url
            return {"title": "Harder, Better, Faster, Stronger", "artist": {"name": "Daft Punk"}}

        with patch.object(bidar, "_http_json", side_effect=fake_json):
            q = bidar._deezer_lookup("https://www.deezer.com/track/3135556")
        self.assertEqual(q, "Daft Punk - Harder, Better, Faster, Stronger")

    def test_youtube_title_query_strips_topic(self):
        with patch.object(bidar, "_http_json",
                          return_value={"title": "Barbie (Remix)", "author_name": "JaidynAlexis - Topic"}):
            q = bidar._youtube_title_query("https://music.youtube.com/watch?v=x")
        self.assertEqual(q, "JaidynAlexis - Barbie (Remix)")

    def test_youtube_title_author_already_in_title(self):
        with patch.object(bidar, "_http_json",
                          return_value={"title": "Joost - Europapa", "author_name": "Joost"}):
            q = bidar._youtube_title_query("https://youtu.be/x")
        self.assertEqual(q, "Joost - Europapa")

    def test_fetch_drm_uses_apple_api_first(self):
        with patch.object(bidar, "_apple_lookup", return_value="Joost - Europapa") as al:
            q = bidar._fetch_drm_metadata("https://music.apple.com/us/song/europapa/1732041802", "apple")
        self.assertEqual(q, "Joost - Europapa")
        al.assert_called_once()


class TestDownloadSyncVideoFallback(unittest.TestCase):
    def test_mp4_used_when_no_audio_file(self):
        import tempfile as _tf
        tmp = _tf.mkdtemp()
        try:
            with open(os.path.join(tmp, "Song.mp4"), "wb") as f:
                f.write(b"0" * 100)
            fake = MagicMock()
            ydl = MagicMock()
            ydl.extract_info.return_value = {"title": "Song", "uploader": "X",
                                             "duration": 100, "thumbnail": ""}
            fake.YoutubeDL.return_value.__enter__ = MagicMock(return_value=ydl)
            fake.YoutubeDL.return_value.__exit__ = MagicMock(return_value=False)
            with patch.object(bidar, "yt_dlp", fake):
                out = bidar._music_download_sync("https://youtube.com/watch?v=x", tmp)
            self.assertIsNotNone(out)
            self.assertTrue(out["filepath"].endswith(".mp4"))
        finally:
            import shutil as _sh
            _sh.rmtree(tmp, ignore_errors=True)

    def test_part_files_ignored(self):
        import tempfile as _tf
        tmp = _tf.mkdtemp()
        try:
            with open(os.path.join(tmp, "Song.mp3.part"), "wb") as f:
                f.write(b"0" * 100)
            fake = MagicMock()
            ydl = MagicMock()
            ydl.extract_info.return_value = {"title": "Song"}
            fake.YoutubeDL.return_value.__enter__ = MagicMock(return_value=ydl)
            fake.YoutubeDL.return_value.__exit__ = MagicMock(return_value=False)
            with patch.object(bidar, "yt_dlp", fake):
                out = bidar._music_download_sync("https://soundcloud.com/x/y", tmp)
            self.assertIsNone(out)
        finally:
            import shutil as _sh
            _sh.rmtree(tmp, ignore_errors=True)


# ────────────────────────────────────────────────────────────────────
# 8. v1.9.4: image aspect-ratio helpers
# ────────────────────────────────────────────────────────────────────
class TestAspectRatio(unittest.TestCase):
    def test_canonical_ratios(self):
        for ar in ("1:1", "16:9", "9:16", "4:3", "3:4", "21:9",
                   "3:2", "2:3", "5:4", "4:5", "1:4", "4:1", "1:8", "8:1"):
            self.assertEqual(bidar._parse_aspect_ratio(ar), ar, ar)

    def test_english_aliases(self):
        self.assertEqual(bidar._parse_aspect_ratio("square"), "1:1")
        self.assertEqual(bidar._parse_aspect_ratio("LANDSCAPE"), "16:9")
        self.assertEqual(bidar._parse_aspect_ratio("portrait"), "9:16")
        self.assertEqual(bidar._parse_aspect_ratio("story"), "9:16")
        self.assertEqual(bidar._parse_aspect_ratio("cinematic"), "21:9")
        self.assertEqual(bidar._parse_aspect_ratio("photo"), "3:2")
        self.assertEqual(bidar._parse_aspect_ratio("tv"), "4:3")

    def test_persian_aliases(self):
        self.assertEqual(bidar._parse_aspect_ratio("مربعی"), "1:1")
        self.assertEqual(bidar._parse_aspect_ratio("افقی"), "16:9")
        self.assertEqual(bidar._parse_aspect_ratio("عمودی"), "9:16")
        self.assertEqual(bidar._parse_aspect_ratio("استوری"), "9:16")
        self.assertEqual(bidar._parse_aspect_ratio("سینمایی"), "21:9")
        self.assertEqual(bidar._parse_aspect_ratio("عکس"), "3:2")

    def test_tolerant_separators(self):
        self.assertEqual(bidar._parse_aspect_ratio("16x9"), "16:9")
        self.assertEqual(bidar._parse_aspect_ratio("16X9"), "16:9")
        self.assertEqual(bidar._parse_aspect_ratio("16×9"), "16:9")
        self.assertEqual(bidar._parse_aspect_ratio(" 9 / 16 "), "9:16")

    def test_invalid(self):
        self.assertIsNone(bidar._parse_aspect_ratio(""))
        self.assertIsNone(bidar._parse_aspect_ratio(None))
        self.assertIsNone(bidar._parse_aspect_ratio("100:1"))
        self.assertIsNone(bidar._parse_aspect_ratio("garbage"))
        self.assertIsNone(bidar._parse_aspect_ratio("0:0"))
        # `4:5` IS supported by Gemini
        self.assertEqual(bidar._parse_aspect_ratio("4:5"), "4:5")

    def test_extract_flag_ar(self):
        ar, p = bidar._extract_ar_flag("--ar 16:9 a red apple on white")
        self.assertEqual(ar, "16:9")
        self.assertEqual(p, "a red apple on white")

    def test_extract_flag_shorthand(self):
        ar, p = bidar._extract_ar_flag("a moody portrait --9:16")
        self.assertEqual(ar, "9:16")
        self.assertEqual(p, "a moody portrait")

    def test_extract_flag_alias(self):
        ar, p = bidar._extract_ar_flag("--landscape sunset over mountains")
        self.assertEqual(ar, "16:9")
        self.assertEqual(p, "sunset over mountains")

    def test_extract_flag_persian_alias(self):
        ar, p = bidar._extract_ar_flag("--استوری یه نقاشی مینیمال از کوه")
        self.assertEqual(ar, "9:16")
        self.assertEqual(p, "یه نقاشی مینیمال از کوه")

    def test_extract_no_flag(self):
        ar, p = bidar._extract_ar_flag("a red apple on white")
        self.assertIsNone(ar)
        self.assertEqual(p, "a red apple on white")

    def test_default_in_config(self):
        self.assertEqual(bidar._DEFAULT_CONFIG["image_aspect_ratio"], "1:1")


class TestGenerateImagePassesAR(unittest.IsolatedAsyncioTestCase):
    async def test_uses_config_default(self):
        bidar.config = copy.deepcopy(bidar._DEFAULT_CONFIG)
        bidar.config["image_aspect_ratio"] = "16:9"
        bidar.config["ai_enabled"] = True
        # Capture the with_params call
        captured = {}

        class FakeChat:
            def with_model(self, *a, **kw): return self
            def with_params(self, **kw):
                captured.update(kw)
                return self
            async def send_message_multimodal_response(self, *a, **kw):
                return ("", [{"data": __import__("base64").b64encode(b"PNGdata").decode()}])

        with patch.object(bidar, "_ai_ready", return_value=(True, "")), \
             patch.object(bidar, "LlmChat", return_value=FakeChat()):
            out, err = await bidar._generate_image("a cat")
        self.assertIsNotNone(out)
        self.assertIsNone(err)
        self.assertEqual(captured.get("image_config"), {"aspect_ratio": "16:9"})

    async def test_explicit_override_wins(self):
        bidar.config = copy.deepcopy(bidar._DEFAULT_CONFIG)
        bidar.config["image_aspect_ratio"] = "1:1"
        bidar.config["ai_enabled"] = True
        captured = {}

        class FakeChat:
            def with_model(self, *a, **kw): return self
            def with_params(self, **kw):
                captured.update(kw)
                return self
            async def send_message_multimodal_response(self, *a, **kw):
                return ("", [{"data": __import__("base64").b64encode(b"PNG").decode()}])

        with patch.object(bidar, "_ai_ready", return_value=(True, "")), \
             patch.object(bidar, "LlmChat", return_value=FakeChat()):
            out, err = await bidar._generate_image("a dog", aspect_ratio="9:16")
        self.assertIsNotNone(out)
        self.assertIsNone(err)
        self.assertEqual(captured.get("image_config"), {"aspect_ratio": "9:16"})


class TestTranslateAutoDetect(unittest.TestCase):
    """`.tl` smart language detection + target resolution."""

    def test_is_persian_text(self):
        self.assertTrue(bidar._is_persian_text("سلام دنیا"))
        self.assertTrue(bidar._is_persian_text("سلام دنیا hello"))  # persian majority
        self.assertFalse(bidar._is_persian_text("hello world"))
        self.assertFalse(bidar._is_persian_text("hola amigo"))
        self.assertFalse(bidar._is_persian_text("hello سلام"))  # latin majority
        self.assertFalse(bidar._is_persian_text(""))
        self.assertFalse(bidar._is_persian_text("12345"))  # no letters at all

    def test_resolve_target_lang_persian_names(self):
        self.assertEqual(bidar._resolve_target_lang("عربی hello there"),
                         ("ar", "hello there"))
        self.assertEqual(bidar._resolve_target_lang("اسپانیایی سلام"),
                         ("es", "سلام"))
        self.assertEqual(bidar._resolve_target_lang("ژاپنی text"),
                         ("ja", "text"))
        self.assertEqual(bidar._resolve_target_lang("کردی"),
                         ("ku", ""))

    def test_resolve_target_lang_english_names(self):
        self.assertEqual(bidar._resolve_target_lang("arabic hello"),
                         ("ar", "hello"))
        self.assertEqual(bidar._resolve_target_lang("ARABIC text"),
                         ("ar", "text"))
        self.assertEqual(bidar._resolve_target_lang("french bonjour"),
                         ("fr", "bonjour"))

    def test_resolve_target_lang_iso_codes(self):
        self.assertEqual(bidar._resolve_target_lang("es hello"),
                         ("es", "hello"))
        self.assertEqual(bidar._resolve_target_lang("ja text"),
                         ("ja", "text"))

    def test_resolve_target_lang_multiword(self):
        self.assertEqual(bidar._resolve_target_lang("ترکی استانبولی merhaba"),
                         ("tr", "merhaba"))
        self.assertEqual(bidar._resolve_target_lang("ترکی آذری salam"),
                         ("az", "salam"))

    def test_resolve_target_lang_no_match(self):
        # Plain text with no leading language token → return (None, original)
        self.assertEqual(bidar._resolve_target_lang("hello world"),
                         (None, "hello world"))
        self.assertEqual(bidar._resolve_target_lang("سلام دنیا"),
                         (None, "سلام دنیا"))

    def test_resolve_target_lang_strips_trailing_punct(self):
        # `.tl عربی,` should still resolve `عربی`
        self.assertEqual(bidar._resolve_target_lang("عربی, سلام"),
                         ("ar", "سلام"))


class TestTranslateAutoDirection(unittest.IsolatedAsyncioTestCase):
    """End-to-end: `.tl <text>` without explicit lang picks fa↔en correctly."""

    async def test_persian_input_translated_to_english(self):
        captured = {}

        async def fake_translate(text, target):
            captured["text"] = text
            captured["target"] = target
            return "Hello world"

        from unittest.mock import AsyncMock
        ev = MagicMock()
        ev.is_reply = False
        ev.sender_id = 12345
        ev.pattern_match = MagicMock()
        ev.pattern_match.group = MagicMock(return_value="سلام دنیا")
        ev.edit = AsyncMock(return_value=MagicMock(edit=AsyncMock()))
        bidar.config = copy.deepcopy(bidar._DEFAULT_CONFIG)
        bidar.config["ai_enabled"] = True

        with patch.object(bidar, "_translate_text", side_effect=fake_translate), \
             patch.object(bidar, "_ai_ready", return_value=(True, "")), \
             patch.object(bidar, "_is_owner", return_value=True):
            await bidar.cmd_translate(ev)
        self.assertEqual(captured["text"], "سلام دنیا")
        self.assertEqual(captured["target"], "en")

    async def test_english_input_translated_to_persian(self):
        captured = {}

        async def fake_translate(text, target):
            captured["target"] = target
            return "سلام"

        from unittest.mock import AsyncMock
        ev = MagicMock()
        ev.is_reply = False
        ev.sender_id = 12345
        ev.pattern_match = MagicMock()
        ev.pattern_match.group = MagicMock(return_value="hello world")
        ev.edit = AsyncMock(return_value=MagicMock(edit=AsyncMock()))

        with patch.object(bidar, "_translate_text", side_effect=fake_translate), \
             patch.object(bidar, "_ai_ready", return_value=(True, "")), \
             patch.object(bidar, "_is_owner", return_value=True):
            await bidar.cmd_translate(ev)
        self.assertEqual(captured["target"], "fa")

    async def test_explicit_lang_overrides_auto(self):
        """`.tl عربی hello` must translate to Arabic, not auto-detect to Persian."""
        captured = {}

        async def fake_translate(text, target):
            captured["text"] = text
            captured["target"] = target
            return "مرحبا"

        from unittest.mock import AsyncMock
        ev = MagicMock()
        ev.is_reply = False
        ev.sender_id = 12345
        ev.pattern_match = MagicMock()
        ev.pattern_match.group = MagicMock(return_value="عربی hello there")
        ev.edit = AsyncMock(return_value=MagicMock(edit=AsyncMock()))

        with patch.object(bidar, "_translate_text", side_effect=fake_translate), \
             patch.object(bidar, "_ai_ready", return_value=(True, "")), \
             patch.object(bidar, "_is_owner", return_value=True):
            await bidar.cmd_translate(ev)
        self.assertEqual(captured["target"], "ar")
        self.assertEqual(captured["text"], "hello there")


class TestImageBytesDetection(unittest.TestCase):
    """Magic-byte detection used by `.r` vision support."""

    def test_jpeg_detected(self):
        self.assertTrue(bidar._is_supported_image_bytes(b"\xff\xd8\xff\xe0" + b"\x00" * 20))
        self.assertTrue(bidar._is_supported_image_bytes(b"\xff\xd8\xff\xe1" + b"\x00" * 20))
        self.assertTrue(bidar._is_supported_image_bytes(b"\xff\xd8\xff\xdb" + b"\x00" * 20))

    def test_png_detected(self):
        self.assertTrue(bidar._is_supported_image_bytes(b"\x89PNG\r\n\x1a\n" + b"\x00" * 20))

    def test_gif_detected(self):
        self.assertTrue(bidar._is_supported_image_bytes(b"GIF87a" + b"\x00" * 20))
        self.assertTrue(bidar._is_supported_image_bytes(b"GIF89a" + b"\x00" * 20))

    def test_webp_detected(self):
        self.assertTrue(bidar._is_supported_image_bytes(b"RIFF\x00\x00\x00\x00WEBP" + b"\x00" * 20))

    def test_unsupported_formats(self):
        # Animated sticker .tgs — must NOT be sent to vision model
        self.assertFalse(bidar._is_supported_image_bytes(b"\x1f\x8b" + b"\x00" * 20))  # gzip header (.tgs)
        # WEBM video sticker
        self.assertFalse(bidar._is_supported_image_bytes(b"\x1aE\xdf\xa3" + b"\x00" * 20))
        # MP4 / random docs
        self.assertFalse(bidar._is_supported_image_bytes(b"\x00\x00\x00\x18ftypmp42" + b"\x00" * 20))
        self.assertFalse(bidar._is_supported_image_bytes(b"PK\x03\x04" + b"\x00" * 20))  # zip
        # Too small / empty
        self.assertFalse(bidar._is_supported_image_bytes(b""))
        self.assertFalse(bidar._is_supported_image_bytes(b"\xff\xd8"))  # 2 bytes only


class TestStyleAgedHelpers(unittest.TestCase):
    def test_resolve_style_english_presets(self):
        for k in ("vangogh", "anime", "ghibli", "pixar", "cyberpunk", "lego", "noir"):
            key, desc = bidar._resolve_style(k)
            self.assertEqual(key, k)
            self.assertTrue(desc and len(desc) > 10)

    def test_resolve_style_case_insensitive(self):
        key, _ = bidar._resolve_style("VANGOGH")
        self.assertEqual(key, "vangogh")
        key, _ = bidar._resolve_style("Pixar")
        self.assertEqual(key, "pixar")

    def test_resolve_style_persian_aliases(self):
        self.assertEqual(bidar._resolve_style("انیمه")[0], "anime")
        self.assertEqual(bidar._resolve_style("گیبلی")[0], "ghibli")
        self.assertEqual(bidar._resolve_style("پیکسار")[0], "pixar")
        self.assertEqual(bidar._resolve_style("ون‌گوگ")[0], "vangogh")
        self.assertEqual(bidar._resolve_style("سایبرپانک")[0], "cyberpunk")
        self.assertEqual(bidar._resolve_style("نوآر")[0], "noir")
        self.assertEqual(bidar._resolve_style("لگو")[0], "lego")

    def test_resolve_style_freeform(self):
        """Unknown style text is passed through as a free-form description."""
        key, desc = bidar._resolve_style("steampunk illustration with brass gears")
        self.assertIsNone(key)
        self.assertEqual(desc, "steampunk illustration with brass gears")

    def test_resolve_style_empty(self):
        self.assertEqual(bidar._resolve_style(""), (None, None))
        self.assertEqual(bidar._resolve_style(None), (None, None))

    def test_parse_age_delta_positive(self):
        self.assertEqual(bidar._parse_age_delta("+20"), 20)
        self.assertEqual(bidar._parse_age_delta("20"), 20)
        self.assertEqual(bidar._parse_age_delta(" 20 "), 20)
        self.assertEqual(bidar._parse_age_delta("20y"), 20)
        self.assertEqual(bidar._parse_age_delta("20 years"), 20)
        self.assertEqual(bidar._parse_age_delta("20 سال"), 20)

    def test_parse_age_delta_negative(self):
        self.assertEqual(bidar._parse_age_delta("-10"), -10)
        self.assertEqual(bidar._parse_age_delta("-5y"), -5)

    def test_parse_age_delta_persian_digits(self):
        self.assertEqual(bidar._parse_age_delta("۲۰"), 20)
        self.assertEqual(bidar._parse_age_delta("-۱۵"), -15)
        self.assertEqual(bidar._parse_age_delta("۱۵ سال"), 15)

    def test_parse_age_delta_invalid(self):
        self.assertIsNone(bidar._parse_age_delta(""))
        self.assertIsNone(bidar._parse_age_delta(None))
        self.assertIsNone(bidar._parse_age_delta("abc"))
        self.assertIsNone(bidar._parse_age_delta("0"))      # zero is meaningless
        self.assertIsNone(bidar._parse_age_delta("100"))    # over cap
        self.assertIsNone(bidar._parse_age_delta("+200y"))  # over cap


class TestSoundCloudDRMFallback(unittest.IsolatedAsyncioTestCase):
    """Picking a SoundCloud search result that's DRM-protected must auto-roll
    over to the next non-DRM result instead of dead-ending the user."""

    async def test_picking_drm_track_rolls_to_next(self):
        bidar.config = copy.deepcopy(bidar._DEFAULT_CONFIG)
        bidar.YTDLP_OK = True
        attempts: list[str] = []

        def fake_dl(target, tmpdir):
            attempts.append(target)
            if "drm-track" in target:
                raise RuntimeError(
                    "ERROR: [soundcloud] 123: This video is DRM protected"
                )
            path = os.path.join(tmpdir, "song.mp3")
            with open(path, "wb") as f:
                f.write(b"FAKE_MP3")
            return {"filepath": path, "title": "Vigen - Chera",
                    "uploader": "Behtarin", "duration": 290, "thumb": None}

        ev = MagicMock()
        ev.out = True
        ev.chat_id = 1
        ev.reply_to_msg_id = None
        ev.message = MagicMock(); ev.message.id = 5
        status = MagicMock()
        status.edit = AsyncMock(); status.delete = AsyncMock()
        ev.edit = AsyncMock(return_value=status)

        sent_file = {}

        async def fake_send_file(chat_id, path, **kw):
            sent_file["path"] = path

        with patch.object(bidar, "_music_download_sync", side_effect=fake_dl), \
             patch.object(bidar, "client") as mc:
            mc.send_file = AsyncMock(side_effect=fake_send_file)
            ok = await bidar._music_download_and_send(
                ev,
                "https://soundcloud.com/drm-track/123",
                "soundcloud",
                False,
                fallback_urls=[
                    "https://soundcloud.com/good-track/2",
                    "https://soundcloud.com/good-track/3",
                ],
            )
        self.assertTrue(ok)
        self.assertEqual(len(attempts), 2)
        self.assertIn("drm-track", attempts[0])
        self.assertIn("good-track/2", attempts[1])

    async def test_all_drm_returns_friendly_message(self):
        bidar.config = copy.deepcopy(bidar._DEFAULT_CONFIG)
        bidar.YTDLP_OK = True

        def all_drm(target, tmpdir):
            raise RuntimeError(
                "ERROR: [soundcloud] X: This video is DRM protected"
            )

        ev = MagicMock()
        ev.out = True; ev.chat_id = 1; ev.reply_to_msg_id = None
        ev.message = MagicMock(); ev.message.id = 7
        last_edit = {}

        async def capture_edit(text, *a, **kw):
            last_edit["text"] = text
        status = MagicMock(); status.edit = AsyncMock(side_effect=capture_edit)
        status.delete = AsyncMock()
        ev.edit = AsyncMock(return_value=status)

        with patch.object(bidar, "_music_download_sync", side_effect=all_drm), \
             patch.object(bidar, "client"):
            ok = await bidar._music_download_and_send(
                ev, "https://soundcloud.com/x/1", "soundcloud", False,
                fallback_urls=["https://soundcloud.com/x/2"],
            )
        self.assertFalse(ok)
        self.assertIn("DRM", last_edit["text"])
        self.assertIn("Pick another", last_edit["text"])


class TestSummaryFlow(unittest.IsolatedAsyncioTestCase):
    async def test_fetch_chat_messages_formats_and_skips(self):
        """Empty messages get media-placeholders, own bot commands are skipped,
        results are reversed to chronological order."""
        # Build a fake message iterator (newest → oldest)
        msgs = []

        def mkmsg(text, sid, name, **flags):
            m = MagicMock()
            m.raw_text = text
            m.sender_id = sid
            m.photo = flags.get("photo")
            m.video = flags.get("video")
            m.video_note = None
            m.voice = flags.get("voice")
            m.sticker = flags.get("sticker")
            m.document = None
            sender = MagicMock()
            sender.first_name = name; sender.last_name = None
            sender.title = None; sender.username = None
            async def gs():
                return sender
            m.get_sender = gs
            return m

        msgs = [
            mkmsg("third message", sid=2, name="Alice"),
            mkmsg(".sum 50", sid=99, name="Me"),  # bot command, skip
            mkmsg("", sid=2, name="Alice", photo=True),  # photo placeholder
            mkmsg("first message", sid=99, name="Me"),
        ]

        async def fake_iter(chat_id, limit):
            for m in msgs:
                yield m

        me = MagicMock(); me.id = 99
        ev = MagicMock(); ev.chat_id = 123

        async def fake_get_me():
            return me

        with patch.object(bidar, "client") as mc:
            mc.get_me = AsyncMock(side_effect=fake_get_me)
            mc.iter_messages = fake_iter
            lines, count = await bidar._fetch_chat_messages(ev, 10)

        # Reversed: oldest → newest = first, photo placeholder, third
        # (the .sum command is filtered out)
        self.assertEqual(count, 3)
        self.assertEqual(lines, ["Me: first message", "Alice: [photo]", "Alice: third message"])



    """Picking a SoundCloud search result that's DRM-protected must auto-roll
    over to the next non-DRM result instead of dead-ending the user."""

    async def test_picking_drm_track_rolls_to_next(self):
        bidar.config = copy.deepcopy(bidar._DEFAULT_CONFIG)
        bidar.YTDLP_OK = True
        attempts: list[str] = []

        def fake_dl(target, tmpdir):
            attempts.append(target)
            if "drm-track" in target:
                # yt-dlp raises this for SoundCloud Go+ paid tracks
                raise RuntimeError(
                    "ERROR: [soundcloud] 123: This video is DRM protected"
                )
            # Simulate a successful download of the next candidate
            path = os.path.join(tmpdir, "song.mp3")
            with open(path, "wb") as f:
                f.write(b"FAKE_MP3")
            # `_music_download_sync` calls `os.listdir(tmpdir)` after yt-dlp;
            # we need to also satisfy its info-dict shape. Mock everything via patch.
            return {"filepath": path, "title": "Vigen - Chera",
                    "uploader": "Behtarin", "duration": 290, "thumb": None}

        # Build an event with a status message we can capture
        ev = MagicMock()
        ev.out = True
        ev.chat_id = 1
        ev.reply_to_msg_id = None
        ev.message = MagicMock(); ev.message.id = 5
        status = MagicMock()
        status.edit = AsyncMock(); status.delete = AsyncMock()
        ev.edit = AsyncMock(return_value=status)

        sent_file = {}

        async def fake_send_file(chat_id, path, **kw):
            sent_file["path"] = path
            sent_file["caption"] = kw.get("caption", "")

        with patch.object(bidar, "_music_download_sync", side_effect=fake_dl), \
             patch.object(bidar, "client") as mc:
            mc.send_file = AsyncMock(side_effect=fake_send_file)
            ok = await bidar._music_download_and_send(
                ev,
                "https://soundcloud.com/drm-track/123",
                "soundcloud",
                False,
                fallback_urls=[
                    "https://soundcloud.com/good-track/2",
                    "https://soundcloud.com/good-track/3",
                ],
            )
        self.assertTrue(ok)
        # First attempt = the DRM URL, second attempt = first fallback
        self.assertEqual(len(attempts), 2)
        self.assertIn("drm-track", attempts[0])
        self.assertIn("good-track/2", attempts[1])
        self.assertEqual(sent_file["path"].endswith("song.mp3"), True)

    async def test_all_drm_returns_friendly_message(self):
        bidar.config = copy.deepcopy(bidar._DEFAULT_CONFIG)
        bidar.YTDLP_OK = True

        def all_drm(target, tmpdir):
            raise RuntimeError(
                "ERROR: [soundcloud] X: This video is DRM protected"
            )

        ev = MagicMock()
        ev.out = True; ev.chat_id = 1; ev.reply_to_msg_id = None
        ev.message = MagicMock(); ev.message.id = 7
        status = MagicMock(); status.edit = AsyncMock(); status.delete = AsyncMock()
        last_edit = {}

        async def capture_edit(text, *a, **kw):
            last_edit["text"] = text
        status.edit = AsyncMock(side_effect=capture_edit)
        ev.edit = AsyncMock(return_value=status)

        with patch.object(bidar, "_music_download_sync", side_effect=all_drm), \
             patch.object(bidar, "client"):
            ok = await bidar._music_download_and_send(
                ev, "https://soundcloud.com/x/1", "soundcloud", False,
                fallback_urls=["https://soundcloud.com/x/2"],
            )
        self.assertFalse(ok)
        # The user should see the friendly DRM hint, not the raw yt-dlp error
        self.assertIn("DRM", last_edit["text"])
        self.assertIn("Pick another", last_edit["text"])



    def test_extract_urls_basic(self):
        urls = bidar._extract_urls("check https://example.com and https://github.com/x/y now")
        self.assertEqual(urls, ["https://example.com", "https://github.com/x/y"])

    def test_extract_urls_keeps_balanced_parentheses(self):
        """Wikipedia-style URLs with balanced (parens) must survive intact."""
        text = "see https://en.wikipedia.org/wiki/Python_(programming_language) now"
        self.assertEqual(bidar._extract_urls(text),
                         ["https://en.wikipedia.org/wiki/Python_(programming_language)"])
        # Trailing `)` without matching `(` IS punctuation
        text2 = "see (https://example.com/foo) end"
        self.assertEqual(bidar._extract_urls(text2), ["https://example.com/foo"])

    def test_extract_urls_dedup_and_limit(self):
        text = " ".join([f"https://x.com/{i}" for i in range(10)])
        urls = bidar._extract_urls(text, max_urls=3)
        self.assertEqual(len(urls), 3)

    def test_extract_urls_empty(self):
        self.assertEqual(bidar._extract_urls(""), [])
        self.assertEqual(bidar._extract_urls(None), [])
        self.assertEqual(bidar._extract_urls("no links here"), [])

    def test_classify_github(self):
        self.assertEqual(bidar._classify_url("https://github.com/torvalds/linux"), "github")
        self.assertEqual(bidar._classify_url("https://www.github.com/foo/bar"), "github")

    def test_classify_youtube(self):
        self.assertEqual(bidar._classify_url("https://www.youtube.com/watch?v=abc"), "youtube")
        self.assertEqual(bidar._classify_url("https://youtu.be/abc"), "youtube")

    def test_classify_generic(self):
        self.assertEqual(bidar._classify_url("https://example.com"), "generic")
        self.assertEqual(bidar._classify_url("https://news.ycombinator.com/item?id=1"), "generic")

    def test_strip_html(self):
        s = "<p>Hello <b>world</b> &amp; co.</p><script>alert(1)</script><style>x{}</style>"
        out = bidar._strip_html(s)
        self.assertNotIn("<", out)
        self.assertNotIn("alert", out)
        self.assertNotIn("x{}", out)
        self.assertIn("Hello world & co.", out)


class TestTldrFlow(unittest.IsolatedAsyncioTestCase):
    async def test_summarise_generic_uses_fetched_payload(self):
        captured = {}

        def fake_fetch(url, max_chars=6000):
            return {"url": url, "title": "Test Page", "description": "desc",
                    "body": "body text " * 50}

        async def fake_ai(prompt):
            captured["prompt"] = prompt
            return "**Headline**\n• point 1\n• point 2"

        with patch.object(bidar, "_fetch_page_text", side_effect=fake_fetch), \
             patch.object(bidar, "_ai_summarise", side_effect=fake_ai):
            out = await bidar._tldr_one("https://example.com/article", "fa")
        self.assertIn("Headline", out)
        self.assertIn("Persian", captured["prompt"])
        self.assertIn("example.com/article", captured["prompt"])

    async def test_github_uses_repo_api(self):
        info = {"full_name": "octocat/Hello-World",
                "description": "test repo", "language": "Python",
                "stars": 1234, "forks": 5, "open_issues": 7,
                "homepage": "", "topics": ["demo"], "license": "MIT",
                "readme": "# Hello"}

        async def fake_ai(prompt):
            return "**Project** does X."

        with patch.object(bidar, "_fetch_github_repo", return_value=info), \
             patch.object(bidar, "_ai_summarise", side_effect=fake_ai):
            out = await bidar._tldr_one("https://github.com/octocat/Hello-World", "en")
        self.assertIn("octocat/Hello-World", out)
        self.assertIn("1,234", out)  # star count formatted
        self.assertIn("Project", out)

    async def test_fetch_failure_returns_error(self):
        def boom(url, max_chars=6000):
            raise TimeoutError("connection timed out")

        with patch.object(bidar, "_fetch_page_text", side_effect=boom):
            out = await bidar._tldr_one("https://broken.example.com/x", "en")
        self.assertIn("❌", out)
        self.assertIn("timed out", out)



    def test_budget_error(self):
        en, fa = bidar._friendly_image_error("Budget has been exceeded! Current cost: 8.15")
        self.assertIn("budget", en.lower())
        self.assertIn("اعتبار", fa)

    def test_safety_block(self):
        en, fa = bidar._friendly_image_error("PROHIBITED_CONTENT")
        self.assertIn("safety", en.lower())
        self.assertIn("فیلتر", fa)

    def test_no_images_returned_is_safety(self):
        en, fa = bidar._friendly_image_error("no images returned (likely content blocked by safety filter)")
        self.assertIn("safety", en.lower())
        self.assertIn("فیلتر", fa)

    def test_rate_limit(self):
        en, fa = bidar._friendly_image_error("Rate limit reached, please retry")
        self.assertIn("rate", en.lower())
        self.assertIn("محدودیت", fa)

    def test_invalid_key(self):
        en, fa = bidar._friendly_image_error("invalid_api_key: 401 Unauthorized")
        self.assertIn("EMERGENT_LLM_KEY", en)
        self.assertIn("EMERGENT_LLM_KEY", fa)

    def test_timeout(self):
        en, fa = bidar._friendly_image_error("Request timed out after 60s")
        self.assertIn("timed out", en.lower())
        self.assertIn("Gemini", fa)

    def test_generic_fallback(self):
        en, fa = bidar._friendly_image_error("some weird error message we don't know")
        self.assertIn("some weird error", en)
        self.assertIn("some weird error", fa)


class TestImageGenReturnsError(unittest.IsolatedAsyncioTestCase):
    async def test_safety_filter_returns_error(self):
        """When Gemini returns no images, generator must surface a safety error."""
        bidar.config = copy.deepcopy(bidar._DEFAULT_CONFIG)
        bidar.config["ai_enabled"] = True

        class EmptyChat:
            def with_model(self, *a, **kw): return self
            def with_params(self, **kw): return self
            async def send_message_multimodal_response(self, *a, **kw):
                return ("", [])  # no images!

        with patch.object(bidar, "_ai_ready", return_value=(True, "")), \
             patch.object(bidar, "LlmChat", return_value=EmptyChat()):
            out, err = await bidar._generate_image("a portrait of someone famous")
        self.assertIsNone(out)
        self.assertIn("safety", (err or "").lower())

    async def test_exception_returns_error(self):
        bidar.config = copy.deepcopy(bidar._DEFAULT_CONFIG)
        bidar.config["ai_enabled"] = True

        class FailingChat:
            def with_model(self, *a, **kw): return self
            def with_params(self, **kw): return self
            async def send_message_multimodal_response(self, *a, **kw):
                raise RuntimeError("Budget has been exceeded!")

        with patch.object(bidar, "_ai_ready", return_value=(True, "")), \
             patch.object(bidar, "LlmChat", return_value=FailingChat()):
            out, err = await bidar._generate_image("any prompt")
        self.assertIsNone(out)
        self.assertIn("budget", (err or "").lower())


# ────────────────────────────────────────────────────────────────────
# GitHub repo URL parsing (.tldr) — `.git` suffix must be stripped safely
# ────────────────────────────────────────────────────────────────────
class TestGithubRepoParse(unittest.TestCase):
    def test_repo_names_ending_in_git_chars_are_preserved(self):
        for url, expected in [
            ("https://github.com/user/audit", ("user", "audit")),
            ("https://github.com/user/chat", ("user", "chat")),
            ("https://github.com/user/tig", ("user", "tig")),
            ("https://github.com/user/config.", ("user", "config.")),
        ]:
            self.assertEqual(bidar._parse_github_repo(url), expected, url)

    def test_git_suffix_stripped(self):
        self.assertEqual(bidar._parse_github_repo("https://github.com/user/repo.git"),
                         ("user", "repo"))
        self.assertEqual(bidar._parse_github_repo("https://github.com/user/tig.git"),
                         ("user", "tig"))

    def test_query_and_www_handled(self):
        self.assertEqual(
            bidar._parse_github_repo("https://www.github.com/User/My-Repo?tab=readme"),
            ("User", "My-Repo"))

    def test_non_github_returns_none(self):
        self.assertIsNone(bidar._parse_github_repo("https://gitlab.com/a/b"))
        self.assertIsNone(bidar._parse_github_repo("https://github.com/onlyowner"))


# ────────────────────────────────────────────────────────────────────
# Whitelist robustness — malformed config entries must never crash
# ────────────────────────────────────────────────────────────────────
class TestWhitelistRobustness(unittest.TestCase):
    def setUp(self):
        bidar.config = copy.deepcopy(bidar._DEFAULT_CONFIG)

    def test_garbage_entries_do_not_crash(self):
        garbage = ["oops", None, "12.5", -1001234567890]
        bidar.config["allowed_groups"] = garbage
        self.assertTrue(bidar._is_chat_allowed(-1001234567890))
        self.assertTrue(bidar._is_chat_allowed(1234567890))
        self.assertFalse(bidar._is_chat_allowed(-999))
        self.assertTrue(bidar._whitelist_contains(1234567890, garbage))
        left = bidar._whitelist_without(-1001234567890, garbage)
        self.assertNotIn(-1001234567890, left)

    def test_numeric_strings_still_match(self):
        bidar.config["allowed_groups"] = ["-1001234567890"]
        self.assertTrue(bidar._is_chat_allowed(1234567890))


# ────────────────────────────────────────────────────────────────────
# Outgoing bot-message prefixes — bot's own captions must be ignored
# by the music auto-detect outgoing handler
# ────────────────────────────────────────────────────────────────────
class TestBotMsgPrefixes(unittest.TestCase):
    def test_bot_generated_messages_are_excluded(self):
        samples = [
            "🎵 **Song**\n👤 Artist\n☁️ SoundCloud",
            '🔍 Search: "https://music.youtube.com/watch?v=x"',
            "🔒 Restricted Search: \"query\"",
            "📰 **TL;DR** of https://example.com",
            "🎨 a prompt with https://music.youtube.com/watch?v=x",
            "▶️ [title](https://music.youtube.com/watch?v=x)",
            "🖼 **Edited:** prompt",
            "👴 Aged: **+20 years**",
            "🧒 Cartoon: **pixar**",
        ]
        for msg in samples:
            self.assertTrue(msg.startswith(bidar._BOT_MSG_PREFIXES), msg)

    def test_normal_owner_message_not_excluded(self):
        self.assertFalse(
            "check this https://music.youtube.com/watch?v=x".startswith(
                bidar._BOT_MSG_PREFIXES))


# ────────────────────────────────────────────────────────────────────
# Text-to-Speech (.say / .voice)
# ────────────────────────────────────────────────────────────────────
class TestTTSHelpers(unittest.TestCase):
    def test_extract_voice_flag_short(self):
        v, rest = bidar._extract_voice_flag("-v onyx read this out loud")
        self.assertEqual(v, "onyx")
        self.assertEqual(rest, "read this out loud")

    def test_extract_voice_flag_long(self):
        v, rest = bidar._extract_voice_flag("hello world --voice nova")
        self.assertEqual(v, "nova")
        self.assertEqual(rest, "hello world")

    def test_extract_voice_flag_unknown_voice_ignored(self):
        v, rest = bidar._extract_voice_flag("-v batman say hi")
        self.assertIsNone(v)
        self.assertEqual(rest, "-v batman say hi")

    def test_extract_voice_flag_none(self):
        v, rest = bidar._extract_voice_flag("just some plain text")
        self.assertIsNone(v)
        self.assertEqual(rest, "just some plain text")

    def test_chunk_text_short_single(self):
        self.assertEqual(bidar._chunk_text("hello", 4000), ["hello"])

    def test_chunk_text_splits_on_words(self):
        words = " ".join(["word"] * 2000)  # ~10k chars
        chunks = bidar._chunk_text(words, 4000)
        self.assertGreater(len(chunks), 1)
        self.assertTrue(all(len(c) <= 4000 for c in chunks))
        # No content lost
        self.assertEqual(" ".join(chunks).split(), words.split())

    def test_opus_duration_from_oggs_header(self):
        # Fake ogg page: 'OggS' + version + flags + 8-byte granule (LE)
        granule = 48000 * 7  # 7 seconds at 48kHz
        page = b"OggS" + b"\x00\x00" + granule.to_bytes(8, "little") + b"rest"
        self.assertEqual(bidar._opus_duration(page), 7)

    def test_opus_duration_bad_data(self):
        self.assertEqual(bidar._opus_duration(b"not-ogg"), 0)
        self.assertEqual(bidar._opus_duration(b""), 0)

    def test_friendly_tts_error_budget(self):
        en, fa = bidar._friendly_tts_error("Error: Budget has been exceeded!")
        self.assertIn("budget", en.lower())
        self.assertIn("اعتبار", fa)

    def test_friendly_tts_error_generic(self):
        en, fa = bidar._friendly_tts_error("weird failure\nline2")
        self.assertIn("weird failure", en)
        self.assertNotIn("line2", en)

    def test_tts_voices_and_models_constants(self):
        self.assertIn("nova", bidar.TTS_VOICES)
        self.assertEqual(len(bidar.TTS_VOICES), 9)
        self.assertEqual(set(bidar.TTS_MODELS), {"tts-1", "tts-1-hd"})

    def test_config_defaults_have_tts_keys(self):
        self.assertEqual(bidar._DEFAULT_CONFIG["tts_voice"], "nova")
        self.assertEqual(bidar._DEFAULT_CONFIG["tts_model"], "tts-1-hd")


class TestTTSGenerate(unittest.IsolatedAsyncioTestCase):
    async def test_generate_returns_bytes(self):
        class FakeTTS:
            def __init__(self, api_key, **kw): pass
            async def generate_speech(self, text, model, voice, response_format):
                return b"OggS-fake-audio"

        with patch.object(bidar, "_tts_ready", return_value=(True, "")), \
             patch.object(bidar, "OpenAITextToSpeech", FakeTTS):
            audio, err = await bidar._tts_generate("سلام", "nova", "tts-1-hd")
        self.assertIsNone(err)
        self.assertEqual(audio, b"OggS-fake-audio")

    async def test_generate_not_ready(self):
        with patch.object(bidar, "_tts_ready", return_value=(False, "no key")):
            audio, err = await bidar._tts_generate("hi", "nova", "tts-1")
        self.assertIsNone(audio)
        self.assertEqual(err, "no key")

    async def test_generate_exception_surfaced(self):
        class FailTTS:
            def __init__(self, api_key, **kw): pass
            async def generate_speech(self, **kw):
                raise RuntimeError("Budget has been exceeded!")

        with patch.object(bidar, "_tts_ready", return_value=(True, "")), \
             patch.object(bidar, "OpenAITextToSpeech", FailTTS):
            audio, err = await bidar._tts_generate("hi", "nova", "tts-1")
        self.assertIsNone(audio)
        self.assertIn("budget", (err or "").lower())


# ────────────────────────────────────────────────────────────────────
# URL Uploader (.up)
# ────────────────────────────────────────────────────────────────────
class TestUploaderHelpers(unittest.TestCase):
    def test_human_size(self):
        self.assertEqual(bidar._human_size(0), "0 B")
        self.assertEqual(bidar._human_size(512), "512 B")
        self.assertEqual(bidar._human_size(1024), "1.0 KB")
        self.assertEqual(bidar._human_size(1536), "1.5 KB")
        self.assertEqual(bidar._human_size(5 * 1024 * 1024), "5.0 MB")
        self.assertEqual(bidar._human_size(2 * 1024**3), "2.0 GB")

    def test_progress_bar(self):
        self.assertEqual(bidar._progress_bar(0), "░" * 10)
        self.assertEqual(bidar._progress_bar(100), "█" * 10)
        self.assertEqual(bidar._progress_bar(40), "████░░░░░░")
        # clamps out-of-range
        self.assertEqual(bidar._progress_bar(150), "█" * 10)
        self.assertEqual(bidar._progress_bar(-5), "░" * 10)

    def test_guess_filename_from_content_disposition(self):
        headers = {"Content-Disposition": 'attachment; filename="report final.pdf"'}
        self.assertEqual(
            bidar._guess_upload_filename("https://x.com/a?b=1", headers), "report final.pdf")

    def test_guess_filename_from_url_path(self):
        self.assertEqual(
            bidar._guess_upload_filename("https://x.com/files/song.mp3", {}), "song.mp3")

    def test_guess_filename_utf8_encoded(self):
        headers = {"Content-Disposition": "attachment; filename*=UTF-8''%D9%81%D8%A7%DB%8C%D9%84.zip"}
        self.assertEqual(bidar._guess_upload_filename("https://x.com/d", headers), "فایل.zip")

    def test_guess_filename_sanitizes_and_falls_back(self):
        # no path, no header → 'file' + ext from mime
        name = bidar._guess_upload_filename("https://x.com/", {}, "image/png")
        self.assertEqual(name, "file.png")

    def test_guess_filename_strips_dangerous_chars(self):
        headers = {"Content-Disposition": 'filename="a/b\\c:d.txt"'}
        got = bidar._guess_upload_filename("https://x.com/d", headers)
        self.assertNotIn("/", got)
        self.assertNotIn("\\", got)
        self.assertTrue(got.endswith(".txt"))

    def test_categorize_upload(self):
        self.assertEqual(bidar._categorize_upload("pic.JPG", ""), "image")
        self.assertEqual(bidar._categorize_upload("clip.mp4", ""), "video")
        self.assertEqual(bidar._categorize_upload("track.flac", ""), "audio")
        self.assertEqual(bidar._categorize_upload("archive.zip", ""), "document")
        # mime-based when extension is unknown
        self.assertEqual(bidar._categorize_upload("noext", "video/webm"), "video")
        self.assertEqual(bidar._categorize_upload("noext", "application/pdf"), "document")

    def test_build_caption_contains_details(self):
        info = {"name": "movie.mp4", "size": 5 * 1024 * 1024,
                "mime": "video/mp4", "url": "https://cdn.example.com/movie.mp4"}
        cap = bidar._build_upload_caption(info, "video")
        self.assertIn("movie.mp4", cap)
        self.assertIn("video/mp4", cap)
        self.assertIn("5.0 MB", cap)
        self.assertIn("cdn.example.com", cap)
        self.assertTrue(cap.startswith("🎬"))

    def test_max_upload_size_is_2gb(self):
        self.assertEqual(bidar.MAX_UPLOAD_SIZE, 2 * 1024 * 1024 * 1024)

    def test_uploader_icons_in_bot_prefixes(self):
        # Uploader caption icons must be skipped by the outgoing music auto-detect handler
        for icon in ("📥", "🎬", "📄", "🖼", "🎵"):
            self.assertTrue((icon + " caption text").startswith(bidar._BOT_MSG_PREFIXES), icon)


class TestUploaderDownload(unittest.IsolatedAsyncioTestCase):
    async def test_download_rejects_oversized_via_content_length(self):
        fake_resp = MagicMock()
        fake_resp.headers = {"Content-Length": str(3 * 1024**3), "Content-Type": "application/zip"}
        fake_resp.close = MagicMock()
        status = MagicMock(); status.edit = AsyncMock()
        with patch.object(bidar.urllib.request, "urlopen", return_value=fake_resp):
            info, err = await bidar._download_url_file(
                "https://x.com/big.zip", tempfile.mkdtemp(), status)
        self.assertIsNone(info)
        self.assertIn("2.0 GB", err)

    async def test_download_streams_small_file(self):
        import io
        data = b"HELLO-FILE-CONTENT" * 10
        stream = io.BytesIO(data)
        fake_resp = MagicMock()
        fake_resp.headers = {"Content-Length": str(len(data)), "Content-Type": "text/plain"}
        fake_resp.read = stream.read
        fake_resp.close = MagicMock()
        status = MagicMock(); status.edit = AsyncMock()
        tmp = tempfile.mkdtemp()
        with patch.object(bidar.urllib.request, "urlopen", return_value=fake_resp):
            info, err = await bidar._download_url_file(
                "https://x.com/hello.txt", tmp, status)
        self.assertIsNone(err)
        self.assertEqual(info["name"], "hello.txt")
        self.assertEqual(info["size"], len(data))
        with open(info["path"], "rb") as f:
            self.assertEqual(f.read(), data)


# ────────────────────────────────────────────────────────────────────
# Image commands: the processing message must be deleted after a successful
# send, and a delete failure must NOT mislabel the send as failed.
# ────────────────────────────────────────────────────────────────────
class TestImageProcessingMessageDeleted(unittest.IsolatedAsyncioTestCase):
    def _make_event(self, is_reply=False):
        status = MagicMock()
        status.delete = AsyncMock()
        status.edit = AsyncMock()
        ev = MagicMock()
        ev.sender_id = 999
        ev.out = True
        ev.chat_id = 55
        ev.is_reply = is_reply
        ev.reply_to_msg_id = None
        ev.pattern_match.group.return_value = "a cute cat"
        ev.edit = AsyncMock(return_value=status)
        async def replied():
            r = MagicMock(); r.media = True; r.id = 7
            return r
        ev.get_reply_message = AsyncMock(side_effect=replied)
        return ev, status

    async def test_cmd_image_deletes_processing_message(self):
        bidar.OWNER_ID = 999
        ev, status = self._make_event()
        with patch.object(bidar, "client") as mc, \
             patch.object(bidar, "_generate_image",
                          AsyncMock(return_value=(b"\x89PNG\r\n\x1a\nfake", None))):
            mc.send_file = AsyncMock()
            await bidar.cmd_image(ev)
        status.delete.assert_awaited_once()
        # Successful send must NOT relabel as failure
        for call in status.edit.await_args_list:
            self.assertNotIn("❌", str(call))

    async def test_cmd_image_delete_failure_does_not_relabel_send(self):
        """If msg.delete() raises after a successful send, the send must NOT be
        reported as failed (regression: delete used to be inside the send try)."""
        bidar.OWNER_ID = 999
        ev, status = self._make_event()
        status.delete = AsyncMock(side_effect=RuntimeError("cannot delete"))
        with patch.object(bidar, "client") as mc, \
             patch.object(bidar, "_generate_image",
                          AsyncMock(return_value=(b"\x89PNG\r\n\x1a\nfake", None))):
            mc.send_file = AsyncMock()
            await bidar.cmd_image(ev)  # must not raise
        mc.send_file.assert_awaited_once()
        # The failure edit must never have been triggered by a delete error
        for call in status.edit.await_args_list:
            self.assertNotIn("img_send_failed", str(call))
            self.assertNotIn("cannot delete", str(call))

    async def test_cmd_imgedit_deletes_processing_message(self):
        bidar.OWNER_ID = 999
        ev, status = self._make_event(is_reply=True)
        with patch.object(bidar, "client") as mc, \
             patch.object(bidar, "_edit_image",
                          AsyncMock(return_value=(b"\x89PNG\r\n\x1a\nfake", None))):
            mc.send_file = AsyncMock()
            mc.download_media = AsyncMock(return_value=b"\xff\xd8\xff\xe0jpeg")
            await bidar.cmd_imgedit(ev)
        status.delete.assert_awaited_once()

    async def test_do_image_transform_deletes_processing_message(self):
        bidar.OWNER_ID = 999
        ev, status = self._make_event(is_reply=True)
        with patch.object(bidar, "client") as mc, \
             patch.object(bidar, "_edit_image",
                          AsyncMock(return_value=(b"\x89PNG\r\n\x1a\nfake", None))):
            mc.send_file = AsyncMock()
            mc.download_media = AsyncMock(return_value=b"\xff\xd8\xff\xe0jpeg")
            await bidar._do_image_transform(ev, "make it anime", "processing", "done")
        status.delete.assert_awaited_once()

    async def test_cmd_image_send_failure_is_reported(self):
        """A genuine send_file failure must still surface an error to the user."""
        bidar.OWNER_ID = 999
        ev, status = self._make_event()
        with patch.object(bidar, "client") as mc, \
             patch.object(bidar, "_generate_image",
                          AsyncMock(return_value=(b"\x89PNG\r\n\x1a\nfake", None))):
            mc.send_file = AsyncMock(side_effect=RuntimeError("network down"))
            await bidar.cmd_image(ev)
        status.delete.assert_not_awaited()
        self.assertTrue(status.edit.await_count >= 1)


# ────────────────────────────────────────────────────────────────────
# Combine two images (.mix)
# ────────────────────────────────────────────────────────────────────
class TestMsgHasImage(unittest.TestCase):
    def test_photo_message(self):
        m = MagicMock(); m.photo = object(); m.document = None
        self.assertTrue(bidar._msg_has_image(m))

    def test_image_document(self):
        m = MagicMock(); m.photo = None
        m.document = MagicMock(); m.document.mime_type = "image/png"
        self.assertTrue(bidar._msg_has_image(m))

    def test_non_image_document(self):
        m = MagicMock(); m.photo = None
        m.document = MagicMock(); m.document.mime_type = "application/pdf"
        self.assertFalse(bidar._msg_has_image(m))

    def test_none(self):
        self.assertFalse(bidar._msg_has_image(None))

    def test_plain_text_message(self):
        m = MagicMock(); m.photo = None; m.document = None
        self.assertFalse(bidar._msg_has_image(m))


class TestCombineImages(unittest.IsolatedAsyncioTestCase):
    async def test_combine_returns_bytes(self):
        class FakeChat:
            def with_model(self, *a, **k): return self
            def with_params(self, **k): return self
            async def send_message_multimodal_response(self, msg):
                return "ok", [{"data": base64.b64encode(b"\x89PNG\r\n\x1a\ncombined").decode()}]
        with patch.object(bidar, "_ai_ready", return_value=(True, "")), \
             patch.object(bidar, "LlmChat", return_value=FakeChat()), \
             patch.object(bidar, "ImageContent", lambda b: b):
            out, err = await bidar._combine_images(["aaa", "bbb"], "blend them")
        self.assertIsNone(err)
        self.assertEqual(out, b"\x89PNG\r\n\x1a\ncombined")

    async def test_combine_needs_two_images(self):
        with patch.object(bidar, "_ai_ready", return_value=(True, "")):
            out, err = await bidar._combine_images(["only-one"], "x")
        self.assertIsNone(out)
        self.assertIn("two", (err or "").lower())

    async def test_combine_no_images_returned_is_safety(self):
        class FakeChat:
            def with_model(self, *a, **k): return self
            def with_params(self, **k): return self
            async def send_message_multimodal_response(self, msg):
                return "text only", []
        with patch.object(bidar, "_ai_ready", return_value=(True, "")), \
             patch.object(bidar, "LlmChat", return_value=FakeChat()), \
             patch.object(bidar, "ImageContent", lambda b: b):
            out, err = await bidar._combine_images(["a", "b"], "x")
        self.assertIsNone(out)
        self.assertIn("no images returned", err)


class TestCmdMixFlow(unittest.IsolatedAsyncioTestCase):
    def _event(self, is_reply=True, own_image=False, caption=""):
        status = MagicMock(); status.delete = AsyncMock(); status.edit = AsyncMock()
        ev = MagicMock()
        ev.sender_id = 999; ev.out = True; ev.chat_id = 55
        ev.is_reply = is_reply; ev.reply_to_msg_id = (7 if is_reply else None)
        ev.pattern_match.group.return_value = caption
        ev.edit = AsyncMock(return_value=status)
        # command message: has image only if own_image
        ev.message = MagicMock()
        ev.message.photo = object() if own_image else None
        ev.message.document = None
        return ev, status

    async def test_needs_two_images(self):
        """Reply to a single image, no attached image → prompts for two images."""
        bidar.OWNER_ID = 999
        ev, status = self._event(is_reply=True, own_image=False)
        replied = MagicMock(); replied.photo = object(); replied.document = None
        replied.grouped_id = None; replied.id = 7
        with patch.object(bidar, "_ai_ready", return_value=(True, "")), \
             patch.object(bidar, "client") as mc:
            mc.download_media = AsyncMock(return_value=b"\xff\xd8\xff\xe0jpeg")
            mc.send_file = AsyncMock()
            ev.get_reply_message = AsyncMock(return_value=replied)
            await bidar.cmd_mix(ev)
        # Only 1 image gathered → must show the "need two" help, never call send_file
        mc.send_file.assert_not_called()
        self.assertTrue(any("mix" in str(c) or "🎭" in str(c)
                            for c in [ev.edit.await_args_list, status.edit.await_args_list]))

    async def test_reply_plus_attached_combines(self):
        """Reply to image A + attach image B → combine and send, delete status."""
        bidar.OWNER_ID = 999
        ev, status = self._event(is_reply=True, own_image=True, caption="on a beach")
        replied = MagicMock(); replied.photo = object(); replied.document = None
        replied.grouped_id = None; replied.id = 7
        sent = {}
        async def fake_send_file(chat_id, path, **kw):
            sent["caption"] = kw.get("caption"); sent["reply_to"] = kw.get("reply_to")
        with patch.object(bidar, "_ai_ready", return_value=(True, "")), \
             patch.object(bidar, "_combine_images",
                          AsyncMock(return_value=(b"\x89PNG\r\n\x1a\nout", None))), \
             patch.object(bidar, "client") as mc:
            mc.download_media = AsyncMock(return_value=b"\xff\xd8\xff\xe0jpeg")
            mc.send_file = AsyncMock(side_effect=fake_send_file)
            ev.get_reply_message = AsyncMock(return_value=replied)
            await bidar.cmd_mix(ev)
        mc.send_file.assert_awaited_once()
        status.delete.assert_awaited_once()
        self.assertIn("on a beach", sent["caption"])

    async def test_reply_to_album_combines(self):
        """Reply to a 2-photo album → both album images are used."""
        bidar.OWNER_ID = 999
        ev, status = self._event(is_reply=True, own_image=False)
        a = MagicMock(); a.photo = object(); a.document = None; a.grouped_id = 111; a.id = 10
        b = MagicMock(); b.photo = object(); b.document = None; b.grouped_id = 111; b.id = 11
        replied = a
        combine_calls = {}
        async def fake_combine(imgs, prompt, aspect_ratio=None):
            combine_calls["n"] = len(imgs); combine_calls["prompt"] = prompt
            return b"\x89PNG\r\n\x1a\nout", None
        with patch.object(bidar, "_ai_ready", return_value=(True, "")), \
             patch.object(bidar, "_combine_images", side_effect=fake_combine), \
             patch.object(bidar, "client") as mc:
            ev.get_reply_message = AsyncMock(return_value=replied)
            mc.get_messages = AsyncMock(return_value=[a, b])
            mc.download_media = AsyncMock(return_value=b"\xff\xd8\xff\xe0jpeg")
            mc.send_file = AsyncMock()
            await bidar.cmd_mix(ev)
        self.assertEqual(combine_calls["n"], 2)
        # No prompt → default blend instruction
        self.assertIn("blend", combine_calls["prompt"].lower())
        mc.send_file.assert_awaited_once()


# ────────────────────────────────────────────────────────────────────
# Document Q&A (.ask)
# ────────────────────────────────────────────────────────────────────
class TestDocHelpers(unittest.TestCase):
    def test_doc_is_supported(self):
        self.assertTrue(bidar._doc_is_supported("report.pdf", ""))
        self.assertTrue(bidar._doc_is_supported("x", "application/pdf"))
        self.assertTrue(bidar._doc_is_supported("notes.txt", "text/plain"))
        self.assertTrue(bidar._doc_is_supported("data.csv", ""))
        self.assertTrue(bidar._doc_is_supported("main.py", ""))
        self.assertTrue(bidar._doc_is_supported("config.json", ""))
        self.assertFalse(bidar._doc_is_supported("photo.jpg", "image/jpeg"))
        self.assertFalse(bidar._doc_is_supported("archive.zip", "application/zip"))
        self.assertFalse(bidar._doc_is_supported("clip.mp4", "video/mp4"))

    def test_extract_text_file(self):
        data = "Hello world\nSecond line".encode("utf-8")
        text, err, trunc = bidar._extract_document_text(data, "notes.txt", "text/plain")
        self.assertIsNone(err)
        self.assertFalse(trunc)
        self.assertIn("Second line", text)

    def test_extract_empty(self):
        text, err, trunc = bidar._extract_document_text(b"   \n  ", "empty.txt", "text/plain")
        self.assertIsNone(text)
        self.assertEqual(err, "empty")

    def test_extract_unsupported(self):
        text, err, trunc = bidar._extract_document_text(b"\x00\x01", "a.zip", "application/zip")
        self.assertIsNone(text)
        self.assertEqual(err, "unsupported")

    def test_extract_truncates(self):
        big = ("x" * (bidar.MAX_DOC_CHARS + 500)).encode("utf-8")
        text, err, trunc = bidar._extract_document_text(big, "big.txt", "text/plain")
        self.assertIsNone(err)
        self.assertTrue(trunc)
        self.assertEqual(len(text), bidar.MAX_DOC_CHARS)


class TestAnswerDocument(unittest.IsolatedAsyncioTestCase):
    async def test_answer_uses_doc_and_history(self):
        captured = {}
        class FakeChat:
            def with_model(self, *a, **k): return self
            async def send_message(self, msg):
                captured["text"] = msg.text
                return "It is about cats."
        doc = {"name": "cats.txt", "text": "Cats are great pets.",
               "history": [("hi", "hello")]}
        with patch.object(bidar, "_ai_ready", return_value=(True, "")), \
             patch.object(bidar, "LlmChat", return_value=FakeChat()):
            ans, err = await bidar._answer_document(doc, "What is this about?")
        self.assertIsNone(err)
        self.assertEqual(ans, "It is about cats.")
        self.assertIn("cats.txt", captured["text"])
        self.assertIn("Cats are great pets.", captured["text"])
        self.assertIn("What is this about?", captured["text"])
        self.assertIn("Q: hi", captured["text"])  # history included

    async def test_answer_not_ready(self):
        with patch.object(bidar, "_ai_ready", return_value=(False, "no key")):
            ans, err = await bidar._answer_document({"name": "x", "text": "y", "history": []}, "q")
        self.assertIsNone(ans)
        self.assertEqual(err, "no key")


class TestCmdAskFlow(unittest.IsolatedAsyncioTestCase):
    def setUp(self):
        bidar.OWNER_ID = 999
        bidar._doc_cache.clear()

    def _event(self, chat_id=55, is_reply=False, raw=""):
        status = MagicMock(); status.delete = AsyncMock(); status.edit = AsyncMock()
        ev = MagicMock()
        ev.sender_id = 999; ev.out = True; ev.chat_id = chat_id
        ev.is_reply = is_reply; ev.reply_to_msg_id = (7 if is_reply else None)
        ev.pattern_match.group.return_value = raw
        ev.edit = AsyncMock(return_value=status)
        return ev, status

    def _doc_reply(self, name="notes.txt", mime="text/plain", size=100):
        r = MagicMock(); r.id = 7
        r.file = MagicMock(); r.file.name = name; r.file.mime_type = mime; r.file.size = size
        return r

    async def test_load_and_answer(self):
        ev, status = self._event(is_reply=True, raw="what is this about?")
        replied = self._doc_reply()
        ev.get_reply_message = AsyncMock(return_value=replied)
        with patch.object(bidar, "client") as mc, \
             patch.object(bidar, "_answer_document",
                          AsyncMock(return_value=("It is about testing.", None))):
            mc.download_media = AsyncMock(return_value=b"This document is about testing.")
            mc.send_message = AsyncMock()
            await bidar.cmd_ask(ev)
        # cached and answered
        self.assertIn(55, bidar._doc_cache)
        self.assertEqual(bidar._doc_cache[55]["name"], "notes.txt")
        self.assertEqual(len(bidar._doc_cache[55]["history"]), 1)
        status.edit.assert_awaited()  # answer edited into status

    async def test_followup_uses_cache(self):
        bidar._doc_cache[55] = {"name": "d.txt", "text": "content",
                                "history": [], "truncated": False, "ts": 0}
        ev, status = self._event(is_reply=False, raw="a follow up question")
        with patch.object(bidar, "client") as mc, \
             patch.object(bidar, "_answer_document",
                          AsyncMock(return_value=("answer 2", None))):
            mc.send_message = AsyncMock()
            await bidar.cmd_ask(ev)
        self.assertEqual(len(bidar._doc_cache[55]["history"]), 1)
        status.edit.assert_awaited()

    async def test_reset_clears_cache(self):
        bidar._doc_cache[55] = {"name": "d", "text": "x", "history": []}
        ev, status = self._event(is_reply=False, raw="reset")
        await bidar.cmd_ask(ev)
        self.assertNotIn(55, bidar._doc_cache)

    async def test_unsupported_file(self):
        ev, status = self._event(is_reply=True, raw="")
        replied = self._doc_reply(name="a.zip", mime="application/zip")
        ev.get_reply_message = AsyncMock(return_value=replied)
        await bidar.cmd_ask(ev)
        self.assertNotIn(55, bidar._doc_cache)
        ev.edit.assert_awaited()

    async def test_no_doc_no_cache_shows_usage(self):
        ev, status = self._event(is_reply=False, raw="")
        await bidar.cmd_ask(ev)
        ev.edit.assert_awaited()  # usage shown, no crash


# ────────────────────────────────────────────────────────────────────
# Word (.docx) support + Server status (.server)
# ────────────────────────────────────────────────────────────────────
def _make_docx_bytes(text: str) -> bytes:
    import docx as _docx
    d = _docx.Document()
    d.add_paragraph(text)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


class TestDocxSupport(unittest.TestCase):
    def test_docx_is_supported(self):
        self.assertTrue(bidar._doc_is_supported("resume.docx", ""))
        self.assertTrue(bidar._doc_is_supported(
            "x", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"))

    def test_extract_docx(self):
        data = _make_docx_bytes("Hello from a Word document")
        text, err, trunc = bidar._extract_document_text(data, "a.docx", "")
        self.assertIsNone(err)
        self.assertIn("Hello from a Word document", text)

    def test_extract_docx_by_mime(self):
        data = _make_docx_bytes("Content via mime")
        text, err, trunc = bidar._extract_document_text(
            data, "noext",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
        self.assertIsNone(err)
        self.assertIn("Content via mime", text)


class TestServerStatus(unittest.TestCase):
    def test_fmt_uptime(self):
        self.assertEqual(bidar._fmt_uptime(0), "0m")
        self.assertEqual(bidar._fmt_uptime(90), "1m")
        self.assertEqual(bidar._fmt_uptime(3661), "1h 1m")
        self.assertEqual(bidar._fmt_uptime(90000), "1d 1h")

    def _fake_data(self):
        return {
            "cpu": 42.0, "cores": 4, "ram_pct": 61.0,
            "ram_used": 2 * 1024**3, "ram_total": 4 * 1024**3,
            "disk_pct": 28.0, "disk_used": 14 * 1024**3, "disk_total": 50 * 1024**3,
            "load": (0.42, 0.55, 0.60), "net_sent": 1024**3, "net_recv": 8 * 1024**3,
            "sys_uptime": 273120, "bot_uptime": 18180,
            "os": "Linux 5.15", "host": "vps-01", "py": "3.11.5", "cpu_temp": None,
        }

    def test_format_en(self):
        s = bidar._format_server_status(self._fake_data(), "en")
        self.assertIn("Live Server Status", s)
        self.assertIn("SERVER", s)
        self.assertIn("42%", s)
        self.assertIn("cores", s)
        self.assertIn("Python", s)
        self.assertIn("```", s)  # wrapped in a monospace code block

    def test_format_fa(self):
        s = bidar._format_server_status(self._fake_data(), "fa")
        self.assertIn("وضعیت", s)
        self.assertIn("42%", s)
        self.assertIn("```", s)

    def test_format_lines_aligned(self):
        """Every framed row must be the same visual width (aligned ASCII box)."""
        s = bidar._format_server_status(self._fake_data(), "en")
        framed = [ln for ln in s.splitlines() if ln.startswith("│ ")]
        self.assertTrue(framed)
        widths = {len(ln) for ln in framed}
        self.assertEqual(len(widths), 1, f"rows misaligned: {widths}")

    def test_ascii_bar(self):
        self.assertEqual(bidar._ascii_bar(0, 16), "░" * 16)
        self.assertEqual(bidar._ascii_bar(100, 16), "█" * 16)
        self.assertEqual(len(bidar._ascii_bar(50, 16)), 16)

    def test_format_with_temp(self):
        d = self._fake_data(); d["cpu_temp"] = 55.0
        s = bidar._format_server_status(d, "en")
        self.assertIn("55C", s)


class TestServerGather(unittest.IsolatedAsyncioTestCase):
    async def test_gather_real_metrics(self):
        d = await bidar._gather_server_status()
        for k in ("cpu", "cores", "ram_pct", "disk_pct", "load",
                  "sys_uptime", "bot_uptime", "os", "py", "host"):
            self.assertIn(k, d)
        self.assertGreaterEqual(d["ram_pct"], 0)


# ────────────────────────────────────────────────────────────────────
# Text → file (.file)
# ────────────────────────────────────────────────────────────────────
class TestBuildFilename(unittest.TestCase):
    def test_bare_extension(self):
        self.assertEqual(bidar._build_filename("py"), "file.py")
        self.assertEqual(bidar._build_filename(".json"), "file.json")
        self.assertEqual(bidar._build_filename("TXT"), "file.TXT")

    def test_full_filename(self):
        self.assertEqual(bidar._build_filename("notes.md"), "notes.md")
        self.assertEqual(bidar._build_filename("my.data.json"), "my.data.json")

    def test_sanitizes_path_chars(self):
        got = bidar._build_filename("../../etc/passwd.txt")
        self.assertNotIn("/", got)
        self.assertTrue(got.endswith(".txt"))

    def test_no_extension_defaults_txt(self):
        self.assertTrue(bidar._build_filename("weird!!name").endswith(".txt")
                        or bidar._build_filename("weird!!name") == "file.weirdname")


class TestCmdMkfile(unittest.IsolatedAsyncioTestCase):
    def _event(self, is_reply=False, raw=""):
        status = MagicMock(); status.delete = AsyncMock(); status.edit = AsyncMock()
        ev = MagicMock()
        ev.sender_id = 999; ev.out = True; ev.chat_id = 55
        ev.is_reply = is_reply; ev.reply_to_msg_id = (7 if is_reply else None)
        ev.pattern_match.group.return_value = raw
        ev.edit = AsyncMock(return_value=status)
        return ev, status

    async def test_text_mode(self):
        bidar.OWNER_ID = 999
        ev, status = self._event(raw='py print("hi")')
        sent = {}
        async def fake_send_file(chat_id, path, **kw):
            with open(path, "rb") as f:
                sent["content"] = f.read().decode()
            sent["name"] = os.path.basename(path)
            sent["force_document"] = kw.get("force_document")
            sent["caption"] = kw.get("caption")
        with patch.object(bidar, "client") as mc:
            mc.send_file = AsyncMock(side_effect=fake_send_file)
            await bidar.cmd_mkfile(ev)
        self.assertEqual(sent["name"], "file.py")
        self.assertEqual(sent["content"], 'print("hi")')
        self.assertTrue(sent["force_document"])
        status.delete.assert_awaited_once()

    async def test_reply_mode(self):
        bidar.OWNER_ID = 999
        ev, status = self._event(is_reply=True, raw="txt")
        replied = MagicMock(); replied.id = 7
        replied.raw_text = "content from replied message"
        replied.text = replied.raw_text
        ev.get_reply_message = AsyncMock(return_value=replied)
        sent = {}
        async def fake_send_file(chat_id, path, **kw):
            with open(path, "rb") as f:
                sent["content"] = f.read().decode()
            sent["name"] = os.path.basename(path)
            sent["reply_to"] = kw.get("reply_to")
        with patch.object(bidar, "client") as mc:
            mc.send_file = AsyncMock(side_effect=fake_send_file)
            await bidar.cmd_mkfile(ev)
        self.assertEqual(sent["name"], "file.txt")
        self.assertEqual(sent["content"], "content from replied message")
        self.assertEqual(sent["reply_to"], 7)

    async def test_full_filename_mode(self):
        bidar.OWNER_ID = 999
        ev, status = self._event(raw="config.json {\"a\": 1}")
        sent = {}
        async def fake_send_file(chat_id, path, **kw):
            sent["name"] = os.path.basename(path)
            with open(path, "rb") as f:
                sent["content"] = f.read().decode()
        with patch.object(bidar, "client") as mc:
            mc.send_file = AsyncMock(side_effect=fake_send_file)
            await bidar.cmd_mkfile(ev)
        self.assertEqual(sent["name"], "config.json")
        self.assertEqual(sent["content"], '{"a": 1}')

    async def test_no_text_errors(self):
        bidar.OWNER_ID = 999
        ev, status = self._event(is_reply=False, raw="py")
        with patch.object(bidar, "client") as mc:
            mc.send_file = AsyncMock()
            await bidar.cmd_mkfile(ev)
        mc.send_file.assert_not_called()
        ev.edit.assert_awaited()

    async def test_usage_when_empty(self):
        bidar.OWNER_ID = 999
        ev, status = self._event(raw="")
        with patch.object(bidar, "client") as mc:
            mc.send_file = AsyncMock()
            await bidar.cmd_mkfile(ev)
        mc.send_file.assert_not_called()


# ────────────────────────────────────────────────────────────────────
# Checkout link generator (.cp)
# ────────────────────────────────────────────────────────────────────
# ────────────────────────────────────────────────────────────────────
# FreeCAD Stripe checkout link (.fc) — usable by members in Allow groups
# ────────────────────────────────────────────────────────────────────
class TestCmdCheckout(unittest.IsolatedAsyncioTestCase):
    def setUp(self):
        bidar.OWNER_ID = 999
        bidar.config = copy.deepcopy(bidar._DEFAULT_CONFIG)

    def _event(self, sender=999, is_private=True, chat_id=999):
        status = MagicMock(); status.edit = AsyncMock(); status.delete = AsyncMock()
        ev = MagicMock()
        ev.sender_id = sender; ev.chat_id = chat_id; ev.is_private = is_private
        ev.pattern_match.group.return_value = None
        ev.respond = AsyncMock(return_value=status)
        ev.delete = AsyncMock()
        return ev, status

    async def test_owner_success_shows_link(self):
        ev, status = self._event(sender=999)
        url = "https://checkout.stripe.com/c/pay/cs_live_ABC#xyz"
        with patch.object(bidar, "_fetch_fc_checkout", return_value=url):
            await bidar.cmd_checkout(ev)
        last = str(status.edit.await_args_list[-1])
        self.assertIn("cs_live_ABC", last)
        ev.delete.assert_awaited()  # original ".fc" command removed

    async def test_member_in_allow_group_can_use(self):
        # non-owner sender, but chat is whitelisted → allowed
        bidar.config["allowed_groups"] = [-1001234567890]
        ev, status = self._event(sender=555, is_private=False, chat_id=-1001234567890)
        with patch.object(bidar, "_fetch_fc_checkout",
                          return_value="https://checkout.stripe.com/c/pay/cs_live_X"):
            await bidar.cmd_checkout(ev)
        ev.respond.assert_awaited()  # responded
        self.assertIn("cs_live_X", str(status.edit.await_args_list[-1]))

    async def test_member_in_non_allow_group_ignored(self):
        # non-owner, chat NOT whitelisted → command ignored (no reply)
        bidar.config["allowed_groups"] = []
        ev, status = self._event(sender=555, is_private=False, chat_id=-100999)
        with patch.object(bidar, "_fetch_fc_checkout",
                          return_value="https://checkout.stripe.com/c/pay/cs_live_X"):
            await bidar.cmd_checkout(ev)
        ev.respond.assert_not_awaited()  # ignored entirely

    async def test_failure_reports(self):
        ev, status = self._event(sender=999)
        with patch.object(bidar, "_fetch_fc_checkout", side_effect=RuntimeError("timeout")):
            await bidar.cmd_checkout(ev)
        self.assertIn("timeout", str(status.edit.await_args_list[-1]))

    async def test_no_link_reports_failure(self):
        ev, status = self._event(sender=999)
        with patch.object(bidar, "_fetch_fc_checkout", return_value=None):
            await bidar.cmd_checkout(ev)
        status.edit.assert_awaited()

    def test_fetch_reads_location_header(self):
        loc = "https://checkout.stripe.com/c/pay/cs_live_ZZZ#frag"
        err = MagicMock()
        err.headers = {"Location": loc}
        real_HTTPError = bidar.urllib.error.HTTPError

        class FakeOpener:
            def open(self, req, timeout=None):
                raise real_HTTPError("u", 303, "See Other", err.headers, None)

        with patch.object(bidar.urllib.request, "build_opener", return_value=FakeOpener()):
            got = bidar._fetch_fc_checkout()
        self.assertEqual(got, loc)


# ────────────────────────────────────────────────────────────────────
# Merge .txt files of a chat (.mergetxt)
# ────────────────────────────────────────────────────────────────────
class TestNormalizeChatRef(unittest.TestCase):
    def test_username(self):
        self.assertEqual(bidar._normalize_chat_ref("@mychannel"), "mychannel")
        self.assertEqual(bidar._normalize_chat_ref("mychannel"), "mychannel")

    def test_link(self):
        self.assertEqual(bidar._normalize_chat_ref("https://t.me/mychannel"), "mychannel")
        self.assertEqual(bidar._normalize_chat_ref("t.me/mychannel"), "mychannel")

    def test_numeric_id(self):
        self.assertEqual(bidar._normalize_chat_ref("-1001234567890"), -1001234567890)


class TestCmdMergeTxt(unittest.IsolatedAsyncioTestCase):
    def _event(self, raw=""):
        status = MagicMock(); status.delete = AsyncMock(); status.edit = AsyncMock()
        ev = MagicMock()
        ev.sender_id = 999; ev.out = True; ev.chat_id = 55; ev.reply_to_msg_id = None
        ev.pattern_match.group.return_value = raw
        ev.edit = AsyncMock(return_value=status)
        return ev, status

    def _txt_msg(self, name, content):
        m = MagicMock()
        m.file = MagicMock(); m.file.name = name; m.file.mime_type = "text/plain"
        m._content = content
        return m

    async def _run(self, raw, messages):
        bidar.OWNER_ID = 999
        ev, status = self._event(raw=raw)
        target = MagicMock(); target.title = "My Channel"; target.username = "mychan"
        sent = {}

        async def fake_send_file(chat_id, path, **kw):
            with open(path, "rb") as f:
                sent["content"] = f.read().decode("utf-8")
            sent["name"] = os.path.basename(path)
            sent["caption"] = kw.get("caption")

        async def fake_iter(entity, filter=None, reverse=False):
            for m in messages:
                yield m

        async def fake_download(m, file=None):
            return m._content.encode("utf-8")

        with patch.object(bidar, "client") as mc:
            mc.get_entity = AsyncMock(return_value=target)
            mc.iter_messages = fake_iter
            mc.download_media = AsyncMock(side_effect=fake_download)
            mc.send_file = AsyncMock(side_effect=fake_send_file)
            await bidar.cmd_mergetxt(ev)
        return sent, status

    async def test_merges_txt_files(self):
        msgs = [self._txt_msg("a.txt", "AAA content"),
                self._txt_msg("b.txt", "BBB content")]
        sent, status = await self._run("https://t.me/mychan", msgs)
        self.assertIn("AAA content", sent["content"])
        self.assertIn("BBB content", sent["content"])
        self.assertIn("FILE 1: a.txt", sent["content"])
        self.assertIn("FILE 2: b.txt", sent["content"])
        self.assertTrue(sent["name"].endswith(".txt"))
        self.assertIn("2txt", sent["name"])
        status.delete.assert_awaited_once()

    async def test_skips_non_txt(self):
        pdf = MagicMock(); pdf.file = MagicMock()
        pdf.file.name = "doc.pdf"; pdf.file.mime_type = "application/pdf"
        pdf._content = "should be skipped"
        txt = self._txt_msg("keep.txt", "keep me")
        sent, status = await self._run("@mychan", [pdf, txt])
        self.assertIn("keep me", sent["content"])
        self.assertNotIn("should be skipped", sent["content"])
        self.assertIn("1txt", sent["name"])

    async def test_no_txt_files(self):
        bidar.OWNER_ID = 999
        ev, status = self._event(raw="@empty")
        target = MagicMock(); target.title = "Empty"; target.username = "empty"
        async def fake_iter(entity, filter=None, reverse=False):
            if False:
                yield None
        with patch.object(bidar, "client") as mc:
            mc.get_entity = AsyncMock(return_value=target)
            mc.iter_messages = fake_iter
            mc.send_file = AsyncMock()
            await bidar.cmd_mergetxt(ev)
        mc.send_file.assert_not_called()
        status.edit.assert_awaited()


class TestParseSizeArg(unittest.TestCase):
    def test_units(self):
        self.assertEqual(bidar._parse_size_arg("100mb"), 100 * 1024**2)
        self.assertEqual(bidar._parse_size_arg("1gb"), 1024**3)
        self.assertEqual(bidar._parse_size_arg("250kb"), 250 * 1024)
        self.assertEqual(bidar._parse_size_arg("500"), 500 * 1024**2)  # bare = MB
        self.assertEqual(bidar._parse_size_arg("2m"), 2 * 1024**2)

    def test_invalid(self):
        self.assertIsNone(bidar._parse_size_arg(""))
        self.assertIsNone(bidar._parse_size_arg(None))
        self.assertIsNone(bidar._parse_size_arg("abc"))
        self.assertIsNone(bidar._parse_size_arg("0kb"))          # below min
        self.assertIsNone(bidar._parse_size_arg("5gb"))          # above 2GB max


class TestCmdSplit(unittest.IsolatedAsyncioTestCase):
    def setUp(self):
        bidar.OWNER_ID = 999

    def _event(self, raw, is_reply=True):
        status = MagicMock(); status.edit = AsyncMock(); status.delete = AsyncMock()
        ev = MagicMock()
        ev.sender_id = 999; ev.out = True; ev.chat_id = 55
        ev.is_reply = is_reply
        ev.pattern_match.group.return_value = raw
        ev.edit = AsyncMock(return_value=status)
        return ev, status

    async def test_splits_at_line_boundaries(self):
        # 20 lines of 500 bytes each → ~10 KB; split at 2kb
        lines = "".join(f"{'x'*499}\n" for _ in range(20))  # 500 bytes/line
        src_bytes = lines.encode()
        replied = MagicMock()
        replied.id = 7
        replied.file = MagicMock(); replied.file.name = "dump.txt"; replied.file.size = len(src_bytes)

        ev, status = self._event("2kb")
        ev.get_reply_message = AsyncMock(return_value=replied)

        sent = []

        async def fake_download(msg, file=None, progress_callback=None):
            with open(file, "wb") as f:
                f.write(src_bytes)
            if progress_callback:
                await progress_callback(len(src_bytes), len(src_bytes))

        async def fake_send_file(chat_id, path, **kw):
            with open(path, "rb") as f:
                sent.append({"name": os.path.basename(path), "content": f.read(),
                             "caption": kw.get("caption"), "reply_to": kw.get("reply_to")})

        with patch.object(bidar, "client") as mc:
            mc.download_media = AsyncMock(side_effect=fake_download)
            mc.send_file = AsyncMock(side_effect=fake_send_file)
            await bidar.cmd_split(ev)

        # every part ≤ 2 KB, and joined content equals the original
        self.assertTrue(len(sent) >= 3)
        for p in sent:
            self.assertLessEqual(len(p["content"]), 2 * 1024)
            # no line cut: each part ends with a full line (newline)
            self.assertTrue(p["content"].endswith(b"\n"))
        joined = b"".join(p["content"] for p in sent)
        self.assertEqual(joined, src_bytes)
        # names are part1, part2...
        self.assertEqual(sent[0]["name"], "dump_part1.txt")
        self.assertEqual(sent[1]["name"], "dump_part2.txt")
        self.assertEqual(sent[0]["reply_to"], 7)

    async def test_no_reply(self):
        ev, status = self._event("100mb", is_reply=False)
        with patch.object(bidar, "client") as mc:
            mc.send_file = AsyncMock()
            await bidar.cmd_split(ev)
        mc.send_file.assert_not_called()
        ev.edit.assert_awaited()

    async def test_bad_size(self):
        replied = MagicMock(); replied.id = 1
        replied.file = MagicMock(); replied.file.name = "x.txt"; replied.file.size = 10
        ev, status = self._event("banana")
        ev.get_reply_message = AsyncMock(return_value=replied)
        with patch.object(bidar, "client") as mc:
            mc.send_file = AsyncMock()
            await bidar.cmd_split(ev)
        mc.send_file.assert_not_called()

    async def test_long_line_gets_own_part(self):
        # single line larger than part size → must still be emitted (can't break line)
        src_bytes = ("y" * 5000 + "\n" + "z" * 10 + "\n").encode()
        replied = MagicMock(); replied.id = 3
        replied.file = MagicMock(); replied.file.name = "big.log"; replied.file.size = len(src_bytes)
        ev, status = self._event("2kb")
        ev.get_reply_message = AsyncMock(return_value=replied)
        sent = []

        async def fake_download(msg, file=None, progress_callback=None):
            with open(file, "wb") as f:
                f.write(src_bytes)

        async def fake_send_file(chat_id, path, **kw):
            with open(path, "rb") as f:
                sent.append(f.read())

        with patch.object(bidar, "client") as mc:
            mc.download_media = AsyncMock(side_effect=fake_download)
            mc.send_file = AsyncMock(side_effect=fake_send_file)
            await bidar.cmd_split(ev)
        joined = b"".join(sent)
        self.assertEqual(joined, src_bytes)
        # first part is the oversized single line (5000+1 bytes)
        self.assertEqual(sent[0], ("y" * 5000 + "\n").encode())


if __name__ == "__main__":
    unittest.main(verbosity=2)